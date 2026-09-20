# -*- coding: utf-8 -*-
"""Motor de presentaciones webshooks (data -> deck).
Definís el contenido en un SPEC (dict) y este módulo arma el .pptx branded.
Cambiás los datos, no el diseño. Estilo tomado de theme_webshooks.json (v4).

    from webshooks_pptx import build
    build(SPEC)                 # SPEC["out"] = ruta de salida

Tipos de slide soportados (campo "type"): resumen, alcance, stack, tabla,
equipo, cronograma, inversion, cierre, seccion.  (La portada va en SPEC["cover"].)
"""
import os, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

_THEME = json.load(open(os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                                     "theme_webshooks.json"), encoding="utf-8"))
K = _THEME["colors"]
NAVY=K["navy"].lstrip("#"); NAVY2="3A4664"; BLUE=K["blue"].lstrip("#"); BLUE_D=K["blue_dark"].lstrip("#")
ORANGE=K["orange"].lstrip("#"); GREEN=K["green"].lstrip("#"); LB=K["tint_blue"].lstrip("#")
LG=K["tint_green"].lstrip("#"); LB3="F4F8FE"; GRID=K["border"].lstrip("#")
BODY=K["body"].lstrip("#"); INK=NAVY; GRAY=K["gray"].lstrip("#"); GRAYL=K["gray_light"].lstrip("#"); WHITE="FFFFFF"
FONT=_THEME["fonts"]["title"]

def _c(h): return RGBColor.from_string(h.lstrip("#"))

class Deck:
    def __init__(self):
        self.prs = Presentation(); self.prs.slide_width=Inches(13.333); self.prs.slide_height=Inches(7.5)
        self.blank = self.prs.slide_layouts[6]; self.pg=1
    def slide(self): return self.prs.slides.add_slide(self.blank)

D=None  # deck actual

def rect(s,l,t,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.RECTANGLE,rad=None):
    sp=s.shapes.add_shape(shape,Inches(l),Inches(t),Inches(w),Inches(h))
    try: sp.shadow.inherit=False
    except: pass
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=_c(fill)
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=_c(line); sp.line.width=Pt(lw)
    if rad is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0]=rad
        except: pass
    return sp

def text(s,l,t,w,h,runs,size=11,color=BODY,bold=False,italic=False,align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP,ls=1.0,sa=2):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,Pt(1))
    if isinstance(runs,str): runs=[runs]
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.line_spacing=ls; p.space_after=Pt(sa); p.space_before=Pt(0)
        if isinstance(para,str): para=[(para,{})]
        for seg in para:
            txt,kw=seg; r=p.add_run(); r.text=txt
            r.font.size=Pt(kw.get("size",size)); r.font.name=FONT
            r.font.color.rgb=_c(kw.get("color",color)); r.font.bold=kw.get("bold",bold)
            r.font.italic=kw.get("italic",italic)
    return tb

def header(s,title,subtitle=None,accent=BLUE):
    rect(s,0,0,0.09,7.5,fill=accent)
    text(s,0.55,0.30,4.0,0.3,[[("web",{"color":BLUE,"bold":True}),("shooks",{"color":INK,"bold":True}),(".",{"color":ORANGE,"bold":True})]],size=15)
    text(s,0.55,0.72,11.5,0.6,title,size=25,color=INK,bold=True)
    rect(s,0.57,1.32,0.9,0.045,fill=ORANGE)
    if subtitle: text(s,0.57,1.42,12.2,0.55,subtitle,size=11.5,color=GRAY,ls=1.02)
    rect(s,0.55,7.06,12.23,0.014,fill=GRID)
    text(s,0.55,7.1,4.0,0.3,_THEME["footer"].get("text","webshooks.com"),size=9,color=GRAYL)
    text(s,12.2,7.1,0.6,0.3,str(D.pg),size=9,color=GRAYL,align=PP_ALIGN.RIGHT); D.pg+=1

def card(s,l,t,w,h,head,body,accent=BLUE,bg=LB3,hsize=12,bsize=9.6,head_h=0.4,bls=1.05):
    rect(s,l,t,w,h,fill=bg,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
    rect(s,l,t,0.07,h,fill=accent,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,l+0.2,t+0.1,w-0.35,head_h,head,size=hsize,color=accent,bold=True)
    if body is not None: text(s,l+0.2,t+0.1+head_h,w-0.38,h-head_h-0.2,body,size=bsize,color=BODY,ls=bls,sa=3)

def stat(s,l,t,w,h,num,lab,accent=BLUE):
    rect(s,l,t,w,h,fill=LB,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08)
    rect(s,l,t,w,0.06,fill=accent)
    text(s,l,t+0.16,w,0.7,num,size=30,color=INK,bold=True,align=PP_ALIGN.CENTER)
    text(s,l+0.1,t+0.92,w-0.2,h-0.98,lab,size=10,color=GRAY,align=PP_ALIGN.CENTER,ls=0.98)

def chip(s,l,t,w,txt,accent=BLUE):
    rect(s,l,t,w,0.36,fill=LB,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,l,t+0.05,w,0.28,txt,size=9.5,color=BLUE_D,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def _bul(items,accent=BLUE,size=10.2):
    out=[]
    for it in items:
        if isinstance(it,(list,tuple)):
            out.append([("▸  ",{"color":accent,"bold":True,"size":size}),(it[0]+"  ",{"color":INK,"bold":True,"size":size}),(it[1],{"color":BODY,"size":size})])
        else:
            out.append([("▸  ",{"color":accent,"bold":True,"size":size}),(it,{"color":BODY,"size":size})])
    return out

def table(s,l,t,w,headers,rows,cw=None,fs=9.5,rh=0.44,hrh=0.44,accent=NAVY):
    n=len(headers); cw=[w/n]*n if not cw else [c*(w/sum(cw)) for c in cw]
    x=l
    for j,hd in enumerate(headers):
        rect(s,x,t,cw[j],hrh,fill=accent); text(s,x+0.08,t,cw[j]-0.12,hrh,hd,size=fs+0.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,ls=0.9); x+=cw[j]
    y=t+hrh
    for i,row in enumerate(rows):
        rect(s,l,y,w,rh,fill=(LB3 if i%2==0 else WHITE),line=GRID,lw=0.5); x=l
        for j,cell in enumerate(row):
            col=INK if j==0 else BODY; bold=(j==0)
            if isinstance(cell,tuple): cell,col=cell; bold=True
            text(s,x+0.08,y,cw[j]-0.12,rh,str(cell),size=fs,color=col,bold=bold,anchor=MSO_ANCHOR.MIDDLE,ls=0.92); x+=cw[j]
        y+=rh
    return y

# ---------------- slide builders (spec-driven) ----------------
def sl_cover(cv):
    s=D.slide(); rect(s,0,0,13.333,7.5,fill=WHITE); rect(s,0,0,4.9,7.5,fill=NAVY); rect(s,0,0,4.9,0.11,fill=ORANGE)
    rect(s,0.5,0.5,0.2,0.2,fill=BLUE,shape=MSO_SHAPE.DIAMOND)
    text(s,0.8,0.42,4.0,0.5,[[("web",{"color":WHITE,"bold":True}),("shooks",{"color":WHITE,"bold":True}),(".",{"color":ORANGE,"bold":True})]],size=26)
    text(s,0.55,1.12,4.2,0.35,cv.get("brand_tag","Desarrollo de Software · LATAM"),size=12.5,color=BLUE)
    rect(s,0.57,1.62,2.4,0.04,fill=ORANGE)
    text(s,0.55,2.1,4.3,1.6,[[(x,{"color":WHITE,"bold":True})] for x in cv.get("title_left",["Propuesta","Comercial"])],size=40,ls=1.0)
    text(s,0.55,3.95,4.3,1.4,[[(cv.get("subtitle_left",""),{"color":GRAYL,"size":12})]],ls=1.1)
    rect(s,0.55,6.05,3.85,0.014,fill=NAVY2)
    text(s,0.55,6.2,4.0,0.35,[[("Cliente   ",{"color":GRAYL,"size":11}),(cv.get("cliente",""),{"color":WHITE,"size":11,"bold":True})]])
    text(s,0.55,6.6,4.0,0.35,[[("Fecha   ",{"color":GRAYL,"size":11}),(cv.get("fecha",""),{"color":WHITE,"size":11,"bold":True})]])
    hl=cv.get("headline",["Iniciá tu","transformación digital","con nosotros."])
    text(s,5.35,0.95,7.5,2.1,[[(hl[0],{"color":INK,"bold":True})]]+[[(h,{"color":INK,"bold":True})] for h in hl[1:-1]]+[[(hl[-1],{"color":BLUE,"bold":True})]],size=33,ls=1.04)
    rect(s,5.4,3.35,7.1,0.04,fill=BLUE)
    text(s,5.35,3.6,7.4,1.0,[[("Propuesta de Solución:  ",{"color":INK,"size":13,"bold":True}),(cv.get("solution",""),{"color":BODY,"size":13})]],ls=1.12)
    x=5.4; y=5.12
    for i,ch in enumerate(cv.get("chips",[])):
        w=0.32+len(ch)*0.10
        if x+w>12.95: x=5.4; y+=0.5
        chip(s,x,y,w,ch,accent=[BLUE,ORANGE,GREEN,BLUE_D][i%4]); x+=w+0.16

def sl_resumen(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"))
    stats=d.get("stats",[]); n=len(stats) or 1; W=(12.23-0.21*(n-1))/n
    for i,st in enumerate(stats):
        stat(s,0.55+i*(W+0.21),2.15,W,1.5,st[0],st[1],accent=[BLUE,ORANGE,GREEN,BLUE_D][i%4])
    if d.get("izq") or d.get("der"):
        text(s,0.55,3.95,12,0.35,d.get("incluye_label","¿Qué incluye?"),size=14,color=INK,bold=True)
        if d.get("izq"): text(s,0.7,4.4,6.0,2.4,_bul(d["izq"],BLUE),ls=1.12,sa=6)
        if d.get("der"): text(s,6.9,4.4,6.0,2.4,_bul(d["der"],GREEN),ls=1.12,sa=6)

def sl_alcance(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"))
    cols=d["cols"]; n=len(cols); gap=0.2; w=(12.23-gap*(n-1))/n
    accents=[BLUE,ORANGE,GREEN,BLUE_D]
    for i,cc in enumerate(cols):
        x=0.55+i*(w+gap); acc=accents[i%4]
        rect(s,x,2.0,w,4.7,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
        rect(s,x,2.0,w,0.55,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08); rect(s,x,2.28,w,0.27,fill=acc)
        text(s,x+0.15,2.0,w-0.3,0.55,cc["head"],size=12.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE)
        text(s,x+0.2,2.7,w-0.36,3.9,_bul(cc["items"],acc,size=9.8),ls=1.1,sa=6)

def sl_stack(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"))
    def tbl(x,tit,rows,acc):
        rect(s,x,2.0,5.9,0.5,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06)
        text(s,x+0.2,2.0,5.5,0.5,tit,size=13,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE); y=2.62
        for k,v in rows:
            rect(s,x,y,5.9,0.44,fill=LB3,line=GRID,lw=0.5)
            text(s,x+0.18,y,2.2,0.44,k,size=10,color=GRAY,anchor=MSO_ANCHOR.MIDDLE)
            text(s,x+2.5,y,3.3,0.44,v,size=10,color=INK,bold=True,anchor=MSO_ANCHOR.MIDDLE); y+=0.5
    tbl(0.55,d.get("left_title","BACKEND"),d["backend"],BLUE)
    tbl(6.9,d.get("right_title","FRONTEND"),d["frontend"],GREEN)
    if d.get("nota"): card(s,0.55,6.35,12.23,0.7,d.get("nota_title","Nota"),[[(d["nota"],{"size":10.5})]],accent=ORANGE,bg="FDF0E4",hsize=11.5,head_h=0.32)

def sl_tabla(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"),accent=d.get("accent",BLUE))
    table(s,0.55,2.05,12.23,d["headers"],d["rows"],cw=d.get("cw"),fs=d.get("fs",10),rh=d.get("rh",0.5))

def sl_equipo(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"))
    team=d["personas"]; w=5.95; accents=[BLUE_D,BLUE,GREEN,ORANGE]
    for i,p in enumerate(team[:4]):
        x=0.55+(i%2)*6.28; y=2.0+(i//2)*2.5; acc=accents[i%4]
        rect(s,x,y,w,2.32,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
        rect(s,x,y,0.08,2.32,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
        text(s,x+0.22,y+0.12,w-0.4,0.4,p["name"],size=12.5,color=acc,bold=True)
        text(s,x+0.28,y+0.58,w-0.5,1.65,_bul(p["items"],acc,size=9.4),ls=1.05,sa=3)

def sl_cronograma(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"),accent=GREEN)
    weeks=d["semanas"]; n=len(weeks); w=(12.23-0.14*(n-1))/n; accents=[GREEN,BLUE,BLUE,ORANGE,NAVY2]
    for i,wk in enumerate(weeks):
        x=0.55+i*(w+0.14); acc=wk.get("color",accents[i%len(accents)])
        rect(s,x,2.15,w,3.9,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
        rect(s,x,2.15,w,0.6,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06); rect(s,x,2.45,w,0.3,fill=acc)
        text(s,x+0.1,2.15,w-0.2,0.6,wk["t"],size=10.8,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
        text(s,x+0.16,2.9,w-0.3,3.0,_bul(wk["items"],acc,size=9.2),ls=1.06,sa=5)

def sl_inversion(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"))
    rect(s,0.55,2.05,5.7,3.1,fill=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
    text(s,0.85,2.35,5.1,0.4,d.get("precio_label","Inversión total"),size=13,color=GRAYL)
    text(s,0.85,2.75,5.1,1.0,d["precio"],size=44,color=WHITE,bold=True)
    text(s,0.85,3.8,5.1,0.4,d.get("precio_nota",""),size=11,color=BLUE)
    if d.get("incluye"): text(s,0.85,4.25,5.1,0.9,[[(d["incluye"],{"color":GRAYL,"size":10.5})]],ls=1.1)
    text(s,6.5,2.05,6.3,0.35,d.get("opciones_label","Opciones de pago"),size=14,color=INK,bold=True); y=2.5
    for op in d.get("opciones",[]):
        acc=op.get("color",BLUE)
        rect(s,6.5,y,6.3,0.82,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06)
        rect(s,6.5,y,0.08,0.82,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
        text(s,6.7,y+0.12,3.5,0.3,op["t"],size=12,color=acc,bold=True)
        text(s,6.7,y+0.44,4.2,0.3,op.get("det",""),size=9,color=GRAY)
        text(s,10.2,y,2.5,0.82,op["monto"],size=16,color=INK,bold=True,align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE); y+=0.95

def sl_cierre(d):
    s=D.slide(); header(s,d["title"],d.get("subtitle"))
    pasos=d.get("pasos",[]); n=len(pasos) or 1; w=(12.23-0.12*(n-1))/n; accents=[BLUE,GREEN,ORANGE,BLUE_D]
    for i,p in enumerate(pasos):
        x=0.55+i*(w+0.12); acc=accents[i%4]
        rect(s,x,2.1,w,2.0,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
        text(s,x+0.2,2.25,1.5,0.7,f"{i+1:02d}",size=30,color=acc,bold=True)
        text(s,x+0.2,2.95,w-0.4,0.35,p["t"],size=13,color=INK,bold=True)
        text(s,x+0.2,3.35,w-0.4,0.7,p.get("d",""),size=9.6,color=BODY,ls=1.08)
    if d.get("resumen"):
        rect(s,0.55,4.45,12.23,1.55,fill=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.03)
        for i,(k,v) in enumerate(d["resumen"]):
            x=0.9+(i%3)*4.05; y=4.7+(i//3)*0.62
            text(s,x,y,3.9,0.3,[[(k+"   ",{"color":GRAYL,"size":10}),(v,{"color":WHITE,"size":11,"bold":True})]])
    if d.get("validez"): text(s,0.55,6.35,12.2,0.4,d["validez"],size=10.5,color=GRAY,italic=True,align=PP_ALIGN.CENTER)

def sl_seccion(d):
    s=D.slide(); rect(s,0,0,13.333,7.5,fill=NAVY); rect(s,0,0,0.09,7.5,fill=d.get("accent",BLUE))
    text(s,0.9,2.4,3.0,2.0,d.get("num",""),size=110,color=d.get("accent",BLUE),bold=True)
    text(s,0.95,4.4,10,1.0,d["title"],size=30,color=WHITE,bold=True)
    if d.get("subtitle"): text(s,0.95,5.4,8,1.2,d["subtitle"],size=13,color=GRAYL,ls=1.15)

_DISPATCH={"resumen":sl_resumen,"alcance":sl_alcance,"stack":sl_stack,"tabla":sl_tabla,
           "equipo":sl_equipo,"cronograma":sl_cronograma,"inversion":sl_inversion,
           "cierre":sl_cierre,"seccion":sl_seccion}

def build(spec):
    global D
    D=Deck()
    if spec.get("cover"): sl_cover(spec["cover"])
    for sl in spec.get("slides",[]):
        fn=_DISPATCH.get(sl["type"])
        if fn: fn(sl)
        else: print("tipo desconocido:", sl["type"])
    out=spec.get("out","propuesta_webshooks.pptx")
    D.prs.save(out); print("deck ->", out, "| slides:", len(D.prs.slides._sldIdLst)); return out
