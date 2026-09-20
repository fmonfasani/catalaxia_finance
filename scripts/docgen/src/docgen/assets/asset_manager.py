from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


class AssetManager:
    """Extracts image assets from PowerPoint slides and materializes them in assets/images."""

    def __init__(self, asset_root: str | Path):
        self.asset_root = Path(asset_root)
        self.images_root = self.asset_root / "images"
        self.images_root.mkdir(parents=True, exist_ok=True)

    def export_images(self, pptx_path: str | Path) -> list[dict[str, str]]:
        export: list[dict[str, str]] = []
        presentation = Presentation(Path(pptx_path))
        for slide_idx, slide in enumerate(presentation.slides, start=1):
            for shape_idx, shape in enumerate(slide.shapes, start=1):
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and getattr(shape, "image", None):
                    image_bytes = shape.image.blob
                    image_name = self._safe_name(slide_idx, shape_idx, image_bytes)
                    destination = self.images_root / image_name
                    destination.write_bytes(image_bytes)
                    export.append({"slide": str(slide_idx), "shape": str(shape_idx), "path": str(destination)})
        return export

    @staticmethod
    def _safe_name(slide_idx: int, shape_idx: int, blob: bytes) -> str:
        digest = hashlib.sha1(blob).hexdigest()[:8]
        return f"slide_{slide_idx}_shape_{shape_idx}_{digest}.png"
