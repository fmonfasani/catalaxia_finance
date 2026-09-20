from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

from docgen.parser.markdown_parser import MarkdownBlock, MarkdownDocument
from docgen.renderer.inline import parse_inline


def _rgb(hexs: str) -> RGBColor:
    return RGBColor.from_string(str(hexs).lstrip("#"))


def _shade(element_pr, hexcolor: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), str(hexcolor).lstrip("#"))
    element_pr.append(shd)


class MarkdownRenderer:
    """Renderiza bloques Markdown a Word aplicando el tema de marca (webshooks)."""

    def __init__(self, document: Document, theme: dict):
        self.doc = document
        self.t = theme
        self.font = theme["fonts"]["body"]
        self.mono = theme["fonts"].get("mono", "Consolas")
        self.base = Path(".")

    def render_document(self, markdown_document: MarkdownDocument) -> None:
        try:
            self.base = Path(markdown_document.source).parent
        except Exception:
            self.base = Path(".")
        for block in markdown_document.blocks:
            self.render_block(block)

    # ---- runs con formato inline ----
    def _inline(self, paragraph, text, color=None, size=None, bold=False):
        color = color or self.t["colors"]["body"]
        size = size or self.t["body"]["size"]
        for s in parse_inline(text or ""):
            r = paragraph.add_run(s.text)
            r.font.name = self.mono if s.code else self.font
            r.font.size = Pt(size)
            r.bold = bool(bold or s.bold)
            r.italic = bool(s.italic)
            if s.code:
                r.font.color.rgb = _rgb(self.t["colors"]["navy"])
            elif s.link:
                r.font.color.rgb = _rgb(self.t["colors"]["links"])
                r.underline = True
            else:
                r.font.color.rgb = _rgb(color)
        return paragraph

    def render_block(self, block: MarkdownBlock) -> None:
        m = getattr(self, f"_r_{block.type}", None)
        if m:
            m(block)
        else:
            self._r_paragraph(block)

    def _r_heading(self, block):
        lvl = block.level or 1
        if lvl > 4:
            lvl = 4
        h = self.t["headings"].get(f"h{lvl}", self.t["headings"]["h3"])
        p = self.doc.add_paragraph()
        try:
            p.style = f"Heading {lvl}"           # necesario para que el índice (TOC) lo tome
        except Exception:
            pass
        self._inline(p, block.text, color=h["color"], size=h["size"], bold=h.get("bold", True))

    def _r_paragraph(self, block):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        self._inline(p, block.text)

    def _r_bullet_list(self, block):
        for item in block.items:
            p = self.doc.add_paragraph(style="List Bullet")
            self._inline(p, item)

    def _r_ordered_list(self, block):
        for item in block.items:
            p = self.doc.add_paragraph(style="List Number")
            self._inline(p, item)

    def _r_quote(self, block):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Inches(0.25)
        pf.space_before = Pt(4)
        pf.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        _shade(pPr, self.t["colors"]["surface"])
        # borde izquierdo azul
        pbdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8"); left.set(qn("w:color"), self.t["colors"]["blue"].lstrip("#"))
        pbdr.append(left); pPr.append(pbdr)
        self._inline(p, block.text, color=self.t["colors"]["gray"])

    def _r_table(self, block):
        rows = block.rows
        if not rows:
            return
        ncol = len(rows[0])
        table = self.doc.add_table(rows=0, cols=ncol)
        table.style = "Table Grid"
        table.autofit = True
        # header
        hdr = table.add_row().cells
        for i, cell_text in enumerate(rows[0]):
            _shade(hdr[i]._tc.get_or_add_tcPr(), self.t["tables"]["header_fill"])
            para = hdr[i].paragraphs[0]
            self._inline(para, cell_text, color=self.t["tables"]["header_font"], size=9.5, bold=True)
        # filas
        for ri, row in enumerate(rows[1:]):
            cells = table.add_row().cells
            for i in range(ncol):
                val = row[i] if i < len(row) else ""
                if ri % 2 == 0:
                    _shade(cells[i]._tc.get_or_add_tcPr(), self.t["tables"]["zebra"])
                self._inline(cells[i].paragraphs[0], val, size=9.5)
        self.doc.add_paragraph()

    def _r_code(self, block):
        lines = (block.text or "").split("\n")
        for ln in lines:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(0); pf.space_after = Pt(0); pf.left_indent = Inches(0.12)
            _shade(p._p.get_or_add_pPr(), self.t["colors"]["code_bg"])
            r = p.add_run(ln if ln else " ")
            r.font.name = self.mono
            r.font.size = Pt(9)
            r.font.color.rgb = _rgb(self.t["colors"]["navy"])
        self.doc.add_paragraph()

    def _r_hrule(self, block):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), self.t["colors"]["border"].lstrip("#"))
        pbdr.append(bottom); pPr.append(pbdr)

    def _r_image(self, block):
        src = (block.text or "").strip()
        path = (self.base / src) if not Path(src).is_absolute() else Path(src)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p.add_run().add_picture(str(path), width=Inches(6.0))
        except Exception:
            self._inline(p, f"*(imagen no encontrada: {src})*", color=self.t["colors"]["gray"])
