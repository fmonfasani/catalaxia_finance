# -*- coding: utf-8 -*-
"""Export de TAGS CRUDOS por empresa/periodo para verificacion independiente.
Sin interpretacion: los valores tal cual estan en cnv_estados_v2 (byma_only) y en
ratios/EDGAR (adr), + la formula de cada ratio + el calendario fiscal. Para que el
usuario haga la cuenta y decida el periodo.
"""
import os, sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY="1A2744"; WHITE="FFFFFF"; GRID="D6E2F2"; LB3="F4F8FE"; INK="1A2744"; GRAYL="6B7589"; ANNUAL="E7F7EC"
def Ft(sz=10,b=False,color="3A4458"): return Font(name="Calibri",size=sz,bold=b,color=color)
def fill(c): return PatternFill("solid",fgColor=c)
S=Side(style="thin",color=GRID); BORDER=Border(left=S,right=S,top=S,bottom=S)
CEN=Alignment(horizontal="center",vertical="center"); LEFT=Alignment(horizontal="left",vertical="center")

DB=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","..","data","screener.db")
con=sqlite3.connect(DB); cur=con.cursor()

# tags que entran en los ratios (orden legible)
INCOME=["Revenue","COGS","GrossProfit","OperatingIncome","EBIT","EBITDA","DA",
        "InterestExpense","IngresosFinancieros","RECPAM","PretaxIncome","IncomeTax","NetIncome"]
BALANCE=["Assets","Liabilities","Equity","DebtCurrent","DebtNonCurrent","Cash"]
CF=["CF_Operativo","CF_Inversion","CashFlowNeto"]
PS=["EPS_basico","EPS_diluido"]
KEY=INCOME+BALANCE+CF+PS

fc={cu:fy for cu,fy in cur.execute("select cuit,fy_end_month from fiscal_calendar")}
byma=cur.execute("select cuit,ticker from mapa_entidades where grupo='byma_only' order by ticker").fetchall()
adr=cur.execute("select cuit,ticker from mapa_entidades where grupo='adr' and coalesce(es_primario,1)=1 order by ticker").fetchall()

wb=Workbook()

# ---------- 1) FORMULAS ----------
fs=wb.active; fs.title="Formulas"
for i,w in enumerate([18,34,44],1): fs.column_dimensions[get_column_letter(i)].width=w
fs.cell(1,1,"CÓMO SE CALCULA CADA RATIO (tags usados)").font=Ft(13,True,NAVY)
fs.cell(2,1,"Los tags están en la hoja 'Tags_BYMA'. Ojo: son ACUMULADOS (YTD) — hay que des-acumular.").font=Ft(9,False,GRAYL)
head=["Ratio","Fórmula","Tags (líneas contables)"]
for i,h in enumerate(head,1):
    c=fs.cell(4,i,h); c.fill=fill(NAVY); c.font=Ft(10,True,WHITE); c.alignment=CEN; c.border=BORDER
FORM=[
 ("ROE","NetIncome / Equity","NetIncome (flujo, TTM) · Equity (balance)"),
 ("ROA","NetIncome / Assets","NetIncome (flujo, TTM) · Assets (balance)"),
 ("Margen Neto","NetIncome / Revenue","NetIncome · Revenue (ambos TTM)"),
 ("Margen Op.","OperatingIncome / Revenue","OperatingIncome · Revenue (ambos TTM)"),
 ("Margen Bruto","GrossProfit / Revenue","GrossProfit · Revenue"),
 ("Deuda/EBITDA","(DebtCurrent+DebtNonCurrent) / EBITDA","DebtCurrent · DebtNonCurrent · EBITDA(TTM)"),
 ("EV/EBITDA","(MarketCap + Deuda − Cash) / EBITDA","DebtCurrent · DebtNonCurrent · Cash · EBITDA(TTM) + MarketCap"),
 ("FCF/CE","(CF_Operativo + CF_Inversion) / (Deuda+Equity)","CF_Operativo · CF_Inversion · Deuda · Equity"),
 ("Payout","Dividendos / NetIncome","cnv_dividendos · NetIncome(TTM)"),
 ("EPS","NetIncome / acciones  (o EPS_basico directo)","EPS_basico  (o NetIncome / shares)"),
 ("PER","Precio / EPS  =  MarketCap / NetIncome","MarketCap · NetIncome(TTM)"),
 ("shares","NetIncome / EPS_basico  (por período)","NetIncome · EPS_basico"),
]
r=5
for ra,fo,tg in FORM:
    fs.cell(r,1,ra).font=Ft(9.5,True,INK); fs.cell(r,2,fo).font=Ft(9.5); fs.cell(r,3,tg).font=Ft(9)
    for i in range(1,4): fs.cell(r,i).border=BORDER; fs.cell(r,i).alignment=LEFT
    r+=1
fs.cell(r+1,1,"TTM = suma de los últimos 4 trimestres standalone (des-acumulados).").font=Ft(9,False,GRAYL)
fs.cell(r+2,1,"Des-acumular: standalone(Q) = YTD(Q) − YTD(Q−1); el Q1 (inicio de ejercicio) ya es standalone.").font=Ft(9,False,GRAYL)

# ---------- 2) CALENDARIO ----------
cs=wb.create_sheet("Calendario")
for i,w in enumerate([10,16,16,14],1): cs.column_dimensions[get_column_letter(i)].width=w
for i,h in enumerate(["Ticker","Fin ejercicio (mes)","Último balance","¿Interino?"],1):
    c=cs.cell(1,i,h); c.fill=fill(NAVY); c.font=Ft(9.5,True,WHITE); c.alignment=CEN; c.border=BORDER
r=2
MES={1:"Ene",2:"Feb",3:"Mar",4:"Abr",5:"May",6:"Jun",7:"Jul",8:"Ago",9:"Sep",10:"Oct",11:"Nov",12:"Dic"}
for cuit,tk in byma:
    up=cur.execute("select ultimo_periodo from screener where cuit=?",(cuit,)).fetchone()
    up=up[0] if up else None
    fy=fc.get(cuit); interino = (up and fy and int(up[5:7])!=fy)
    cs.cell(r,1,tk).font=Ft(9.5,True,INK)
    cs.cell(r,2,MES.get(fy,"?") if fy else "—").font=Ft(9.5)
    cs.cell(r,3,str(up)).font=Ft(9.5)
    cc=cs.cell(r,4,"SÍ (ojo)" if interino else "no"); cc.font=Ft(9.5,True,"C0392B" if interino else "1F6B35")
    for i in range(1,5): cs.cell(r,i).border=BORDER; cs.cell(r,i).alignment=CEN if i!=1 else LEFT
    r+=1

# ---------- 3) TAGS BYMA (crudos, largo) ----------
ts=wb.create_sheet("Tags_BYMA")
for i,w in enumerate([9,22,13,7,18],1): ts.column_dimensions[get_column_letter(i)].width=w
for i,h in enumerate(["Ticker","Tag (concepto)","period_end","tipo","valor"],1):
    c=ts.cell(1,i,h); c.fill=fill(NAVY); c.font=Ft(9.5,True,WHITE); c.alignment=CEN; c.border=BORDER
ts.freeze_panes="A2"
r=2
for cuit,tk in byma:
    # ultimos 8 periodos
    pers=[x[0] for x in cur.execute("select distinct period_end from cnv_estados_v2 where cuit=? order by period_end desc limit 8",(cuit,))]
    fy=fc.get(cuit)
    for con2 in KEY:
        rows=cur.execute("select period_end,tipo,valor from cnv_estados_v2 where cuit=? and concepto=? and period_end in (%s) order by period_end desc" % ",".join("?"*len(pers)),
                         (cuit,con2,*pers)).fetchall()
        for pe,tipo,val in rows:
            annual = fy and int(pe[5:7])==fy
            ts.cell(r,1,tk).font=Ft(9)
            ts.cell(r,2,con2).font=Ft(9,color=INK)
            ce=ts.cell(r,3,pe); ce.font=Ft(9);
            if annual: ce.fill=fill(ANNUAL); ce.font=Ft(9,True,"1F6B35")
            ts.cell(r,4,tipo).font=Ft(8.5,color=GRAYL)
            vv=ts.cell(r,5,round(val,2) if isinstance(val,(int,float)) else val); vv.font=Ft(9); vv.alignment=Alignment(horizontal="right")
            for i in range(1,6): ts.cell(r,i).border=BORDER
            r+=1

# ---------- 4) TAGS ADR (EDGAR, ya TTM) ----------
asheet=wb.create_sheet("Tags_ADR")
for i,w in enumerate([9,26,18],1): asheet.column_dimensions[get_column_letter(i)].width=w
for i,h in enumerate(["Ticker","Concepto (EDGAR, TTM/snapshot)","valor"],1):
    c=asheet.cell(1,i,h); c.fill=fill(NAVY); c.font=Ft(9.5,True,WHITE); c.alignment=CEN; c.border=BORDER
rcols=[("_revenue_ttm","Revenue TTM"),("_netincome_ttm","NetIncome TTM"),("_ebitda_ttm","EBITDA TTM"),
       ("_fcf_ttm","FCF TTM"),("_equity","Equity"),("_assets","Assets"),("_deuda","Deuda")]
r=2
for cuit,tk in adr:
    rr=cur.execute("select ticker,%s from ratios where ticker=?" % ",".join(c for c,_ in rcols),(tk,)).fetchone()
    if not rr:
        asheet.cell(r,1,tk).font=Ft(9); asheet.cell(r,2,"(sin fila en ratios EDGAR)").font=Ft(9,color=GRAYL); r+=1; continue
    for j,(col,nm) in enumerate(rcols):
        v=rr[1+j]
        asheet.cell(r,1,tk).font=Ft(9); asheet.cell(r,2,nm).font=Ft(9,color=INK)
        vv=asheet.cell(r,3,round(v,2) if isinstance(v,(int,float)) else v); vv.alignment=Alignment(horizontal="right"); vv.font=Ft(9)
        for i in range(1,4): asheet.cell(r,i).border=BORDER
        r+=1

out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"Tags_Crudos_Catalaxia.xlsx")
wb.save(out)
nb=cur.execute("select count(distinct cuit||concepto) from cnv_estados_v2 where cuit in (%s)" % ",".join("'%s'"%c for c,_ in byma)).fetchone()[0]
print("SAVED:",out)
print("hojas:",wb.sheetnames)
print(f"byma_only: {len(byma)} empresas  |  adr: {len(adr)}  |  tags exportados: {len(KEY)} conceptos x 8 periodos")
con.close()
