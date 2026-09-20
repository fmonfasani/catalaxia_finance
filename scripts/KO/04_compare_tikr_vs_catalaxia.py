#!/usr/bin/env python3
"""
Catalaxia / KO — 04 Comparador puro TIKR vs Catalaxia.
Primary layer: VALIDACIÓN.

Receives two homologated structures (TIKR reference + Catalaxia TIKR-shaped
output) and compares by metric label and period date.

Does NOT calculate ratios, repair mappings, or modify source values.

Statuses:
  MATCH              ≤0.5% or ≤0.005 absolute
  CLOSE              ≤5% or ≤0.05 absolute
  DIFFERENCE         >5%
  CATALAXIA_NULL     Catalaxia has no value
  TIKR_NULL          TIKR has no value
  BOTH_NULL          Neither has a value
"""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

BASE = Path(__file__).resolve().parent

TIKR_CANDIDATES = [
    BASE / "KO_TIKR_REFERENCE_QUARTERLY_MASTER.xlsx",
    BASE / "KO_TIKR_REFERENCE_QUARTERLY.xlsx",
]
CAT = BASE / "KO_CATALAXIA_TIKR_OUTPUT_QUARTERLY.xlsx"
OUT = BASE / "KO_TIKR_VS_CATALAXIA_DIFFERENCES.xlsx"

SHEETS = ["TIKR_IS", "TIKR_BS", "TIKR_CF", "TIKR_RATIOS"]

SECTION_HEADERS = {
    "supplementary data:", "return ratios:", "margin analysis:",
    "asset turnover:", "short term liquidity:", "long-term solvency:",
    "price factors:",
}

PCT_CHANGE_LABELS = {
    "% change yoy", "% gross margins", "% operating margins",
    "% net income to common incl extra items margins",
    "% net income to common excl. extra items margins",
    "% free cash flow margins",
}


def find_tikr() -> Path:
    for p in TIKR_CANDIDATES:
        if p.exists():
            return p
    raise FileNotFoundError(
        f"TIKR reference not found: {[str(p) for p in TIKR_CANDIDATES]}"
    )


def to_num(x) -> float:
    if isinstance(x, (int, float)):
        return float(x)
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return np.nan
    s = str(x).strip()
    s = s.replace(",", "")
    s = re.sub(r"[%x]$", "", s)
    s = s.replace("(", "-").replace(")", "")
    s = s.replace("USD ", "").strip()
    try:
        return float(s)
    except (ValueError, TypeError):
        return np.nan


def is_section_or_derived(label: str) -> bool:
    norm = str(label).strip().lower()
    if norm in SECTION_HEADERS:
        return True
    if norm in PCT_CHANGE_LABELS:
        return True
    if norm.startswith("%"):
        return True
    return pd.isna(label) or str(label).strip() == ""


def compare_sheet(tikr_df: pd.DataFrame, cat_df: pd.DataFrame) -> pd.DataFrame:
    tikr_df = tikr_df.copy()
    cat_df = cat_df.copy()

    tikr_df.columns = [str(c).strip() for c in tikr_df.columns]
    cat_df.columns = [str(c).strip() for c in cat_df.columns]

    t_key = next((c for c in tikr_df.columns if c.lower() == "metric"), None)
    c_key = next((c for c in cat_df.columns if c.lower() == "metric"), None)
    if t_key is None or c_key is None:
        raise ValueError("Missing 'Metric' column")

    tikr_df = tikr_df.rename(columns={t_key: "metric"})
    cat_df = cat_df.rename(columns={c_key: "metric"})

    period_cols = [
        c for c in tikr_df.columns
        if c != "metric" and c in cat_df.columns
    ]

    cat_idx = cat_df.drop_duplicates("metric").set_index("metric")

    rows = []
    for _, tr in tikr_df.iterrows():
        metric = tr["metric"]
        if is_section_or_derived(metric):
            continue

        if metric not in cat_idx.index:
            for pcol in period_cols:
                tv = to_num(tr[pcol])
                rows.append([metric, pcol, tv, np.nan, np.nan, np.nan,
                             "CATALAXIA_NULL" if not np.isnan(tv) else "BOTH_NULL"])
            continue

        cr = cat_idx.loc[metric]
        for pcol in period_cols:
            tv = to_num(tr[pcol])
            cv = to_num(cr[pcol]) if pcol in cr.index else np.nan

            if np.isnan(tv) and np.isnan(cv):
                status = "BOTH_NULL"
                diff = pct_diff = np.nan
            elif np.isnan(tv):
                status = "TIKR_NULL"
                diff = pct_diff = np.nan
            elif np.isnan(cv):
                status = "CATALAXIA_NULL"
                diff = pct_diff = np.nan
            else:
                diff = cv - tv
                abs_tv = abs(tv)
                threshold_match = max(abs_tv * 0.005, 0.005)
                threshold_close = max(abs_tv * 0.05, 0.05)
                pct_diff = (diff / tv * 100) if tv != 0 else np.nan

                if abs(diff) <= threshold_match:
                    status = "MATCH"
                elif abs(diff) <= threshold_close:
                    status = "CLOSE"
                else:
                    status = "DIFFERENCE"

            rows.append([metric, pcol, tv, cv, diff, pct_diff, status])

    return pd.DataFrame(
        rows,
        columns=["metric", "period", "TIKR", "CATALAXIA",
                 "DIFF_CAT_MINUS_TIKR", "DIFF_%", "STATUS"],
    )


def main():
    tikr_path = find_tikr()
    if not CAT.exists():
        raise FileNotFoundError(f"Catalaxia output not found: {CAT}")

    tikr_xls = pd.ExcelFile(tikr_path)
    cat_xls = pd.ExcelFile(CAT)

    summary = {}
    with pd.ExcelWriter(OUT, engine="openpyxl") as writer:
        for sheet in SHEETS:
            diff_sheet = sheet.replace("TIKR_", "DIFF_")
            result = compare_sheet(
                pd.read_excel(tikr_xls, sheet),
                pd.read_excel(cat_xls, sheet),
            )
            result.to_excel(writer, sheet_name=diff_sheet, index=False)
            summary[diff_sheet] = result["STATUS"].value_counts(dropna=False).to_dict()

    wb = load_workbook(OUT)
    fills = {
        "MATCH": "C6EFCE",
        "CLOSE": "FFEB9C",
        "DIFFERENCE": "FFC7CE",
        "TIKR_NULL": "D9E1F2",
        "CATALAXIA_NULL": "FCE4D6",
        "BOTH_NULL": "E7E6E6",
    }
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for row_idx in range(2, ws.max_row + 1):
            st = ws.cell(row_idx, 7).value
            if st in fills:
                ws.cell(row_idx, 7).fill = PatternFill(
                    "solid", fgColor=fills[st]
                )
    wb.save(OUT)

    print("=" * 68)
    print("KO TIKR vs CATALAXIA COMPARISON")
    print("=" * 68)
    print(f"TIKR reference: {tikr_path.name}")
    print(f"Catalaxia output: {CAT.name}")
    print(f"Differences output: {OUT.name}")
    print()

    total_match = total_close = total_diff = 0
    total_cat_null = total_tikr_null = total_both_null = 0

    for sheet_name, counts in summary.items():
        print(f"  {sheet_name}:")
        for status, count in sorted(counts.items()):
            print(f"    {status}: {count}")
        total_match += counts.get("MATCH", 0)
        total_close += counts.get("CLOSE", 0)
        total_diff += counts.get("DIFFERENCE", 0)
        total_cat_null += counts.get("CATALAXIA_NULL", 0)
        total_tikr_null += counts.get("TIKR_NULL", 0)
        total_both_null += counts.get("BOTH_NULL", 0)
        print()

    total = total_match + total_close + total_diff + total_cat_null + total_tikr_null + total_both_null
    comparable = total_match + total_close + total_diff
    print(f"TOTAL COMPARISONS: {total}")
    print(f"  COMPARABLE: {comparable}  (MATCH={total_match}, CLOSE={total_close}, DIFFERENCE={total_diff})")
    print(f"  NULL: {total_cat_null + total_tikr_null + total_both_null}  (CAT_NULL={total_cat_null}, TIKR_NULL={total_tikr_null}, BOTH_NULL={total_both_null})")
    print()

    execution_status = "PASS"
    if comparable > 0:
        match_rate = (total_match + total_close) / comparable * 100
        print(f"MATCH+CLOSE rate: {match_rate:.1f}%")
    else:
        match_rate = 0
        print("WARNING: No comparable metrics found")

    print(f"EXECUTION_STATUS: {execution_status}")
    print("=" * 68)


if __name__ == "__main__":
    main()
