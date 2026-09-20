# -*- coding: utf-8 -*-
"""Propuesta webshooks -> Catalaxia v4.1 — reproducible.

v4.1 agrega 3 slides sobre v4.0:
  · Cobertura de datos por ratio y fuente (EDGAR / ADR / Only-BYMA) — valores reales.
  · Circuito de entregas por sábado (los 108 puntos, paso a paso).
  · Desglose de los 108 puntos por épica (referencias marcadas).

Los % de cobertura salen de data/screener.db (tabla `screener`, columna `grupo`).
Correr:  python docs/proyecto/generar_propuesta_v4.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ---- paleta v3/v4 ----
NAVY="1A2744"; NAVY2="3A4664"; NAVY_D="10192E"; BLUE="3B82F6"; BLUE_D="1A4D8F"
ORANGE="F5821F"; ORANGE_L="F5B97F"; GREEN="43A047"; GREEN_D="1F6B35"
LB="DCEAFF"; LG="E7F7EC"; LB3="F4F8FE"; GRID="D6E2F2"
BODY="3A4458"; INK="1A2744"; GRAY="6B7589"; GRAYL="AAB4C8"; WHITE="FFFFFF"
# tintes para el heatmap
HM_AMBER="FDEBD8"; HM_AMBER_T="B8621B"; HM_RED="FBE4E4"; HM_RED_T="C0392B"
FONT="Calibri"

def C(h): return RGBColor.from_string(h)
prs = Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK = prs.slide_layouts[6]
PG={"n":1}

def slide(): return prs.slides.add_slide(BLANK)
def _noshadow(sp):
    try: sp.shadow.inherit=False
    except: pass

def rect(s,l,t,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.RECTANGLE,rad=None):
    sp=s.shapes.add_shape(shape,Inches(l),Inches(t),Inches(w),Inches(h)); _noshadow(sp)
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=C(fill)
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=C(line); sp.line.width=Pt(lw)
    if rad is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0]=rad
        except: pass
    return sp

def text(s,l,t,w,h,runs,size=11,color=BODY,bold=False,italic=False,align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP,font=FONT,ls=1.0,sa=2,wrap=True):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=wrap; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,Pt(1))
    if isinstance(runs,str): runs=[runs]
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.line_spacing=ls; p.space_after=Pt(sa); p.space_before=Pt(0)
        if isinstance(para,str): para=[(para,{})]
        for seg in para:
            txt,kw=seg; r=p.add_run(); r.text=txt
            r.font.size=Pt(kw.get("size",size)); r.font.name=kw.get("font",font)
            r.font.color.rgb=C(kw.get("color",color)); r.font.bold=kw.get("bold",bold)
            r.font.italic=kw.get("italic",italic)
    return tb

def header(s,title,subtitle=None,accent=BLUE):
    rect(s,0,0,0.09,7.5,fill=accent)
    text(s,0.55,0.30,4.0,0.3,[[("web",{"color":BLUE,"bold":True}),("shooks",{"color":INK,"bold":True}),(".",{"color":ORANGE,"bold":True})]],size=15)
    text(s,0.55,0.72,11.5,0.6,title,size=25,color=INK,bold=True)
    rect(s,0.57,1.32,0.9,0.045,fill=ORANGE)
    if subtitle:
        text(s,0.57,1.42,12.2,0.55,subtitle,size=11.5,color=GRAY,ls=1.02)
    # footer
    rect(s,0.55,7.06,12.23,0.014,fill=GRID)
    text(s,0.55,7.1,4.0,0.3,"webshooks.com",size=9,color=GRAYL)
    text(s,12.2,7.1,0.6,0.3,str(PG["n"]),size=9,color=GRAYL,align=PP_ALIGN.RIGHT); PG["n"]+=1

def card(s,l,t,w,h,head,body,accent=BLUE,bg=LB3,hsize=12,bsize=9.6,head_h=0.4,bls=1.05):
    rect(s,l,t,w,h,fill=bg,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
    rect(s,l,t,0.07,h,fill=accent,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,l+0.2,t+0.1,w-0.35,head_h,head,size=hsize,color=accent,bold=True)
    if body is not None:
        text(s,l+0.2,t+0.1+head_h,w-0.38,h-head_h-0.2,body,size=bsize,color=BODY,ls=bls,sa=3)

def stat(s,l,t,w,h,num,lab,accent=BLUE):
    rect(s,l,t,w,h,fill=LB,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08)
    rect(s,l,t,w,0.06,fill=accent)
    text(s,l,t+0.16,w,0.7,num,size=30,color=INK,bold=True,align=PP_ALIGN.CENTER)
    text(s,l+0.1,t+0.92,w-0.2,h-0.98,lab,size=10,color=GRAY,align=PP_ALIGN.CENTER,ls=0.98)

def chip(s,l,t,w,txt,accent=BLUE):
    rect(s,l,t,w,0.36,fill=LB,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,l,t+0.05,w,0.28,txt,size=9.5,color=BLUE_D,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def table(s,l,t,w,headers,rows,cw=None,fs=9.5,rh=0.44,hrh=0.44,accent=NAVY):
    n=len(headers)
    if cw is None: cw=[w/n]*n
    else:
        sc=w/sum(cw); cw=[c*sc for c in cw]
    x=l
    for j,hd in enumerate(headers):
        rect(s,x,t,cw[j],hrh,fill=accent)
        text(s,x+0.08,t,cw[j]-0.12,hrh,hd,size=fs+0.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,ls=0.9)
        x+=cw[j]
    y=t+hrh
    for i,row in enumerate(rows):
        bg=LB3 if i%2==0 else WHITE
        rect(s,l,y,w,rh,fill=bg,line=GRID,lw=0.5)
        x=l
        for j,cell in enumerate(row):
            col=INK if j==0 else BODY; bold=(j==0)
            if isinstance(cell,tuple): cell,col=cell; bold=True
            text(s,x+0.08,y,cw[j]-0.12,rh,str(cell),size=fs,color=col,bold=bold,anchor=MSO_ANCHOR.MIDDLE,ls=0.92)
            x+=cw[j]
        y+=rh
    return y

def bullets(items, accent=BLUE, size=10.2):
    out=[]
    for it in items:
        if isinstance(it,tuple):
            out.append([("▸  ",{"color":accent,"bold":True,"size":size}),(it[0]+"  ",{"color":INK,"bold":True,"size":size}),(it[1],{"color":BODY,"size":size})])
        else:
            out.append([("▸  ",{"color":accent,"bold":True,"size":size}),(it,{"color":BODY,"size":size})])
    return out

# ---- helpers nuevos v4.1 -------------------------------------------------
def _hm_color(p):
    """Color de celda del heatmap según % de cobertura."""
    if p>=85: return (LG, GREEN_D)
    if p>=70: return (LB, BLUE_D)
    if p>=50: return (HM_AMBER, HM_AMBER_T)
    return (HM_RED, HM_RED_T)

def heatmap(s,l,t,w,col_heads,rows,rowh=0.285,name_w=2.7,hrh=0.44):
    """col_heads: [(label, accent), ...]  |  rows: [(ratio, [(pct,n),(pct,n),...]), ...]"""
    ncol=len(col_heads); cellw=(w-name_w)/ncol
    text(s,l+0.04,t,name_w-0.08,hrh,"Ratio",size=9.8,color=GRAY,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    x=l+name_w
    for lab,acc in col_heads:
        rect(s,x,t,cellw-0.06,hrh,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08)
        text(s,x,t,cellw-0.06,hrh,[[(lab,{"color":WHITE,"size":9.3,"bold":True})]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,ls=0.9)
        x+=cellw
    y=t+hrh+0.05
    for ratio,cells in rows:
        text(s,l+0.04,y,name_w-0.08,rowh,ratio,size=9.4,color=INK,bold=True,anchor=MSO_ANCHOR.MIDDLE)
        x=l+name_w
        for pct,n in cells:
            fill,tcol=_hm_color(pct)
            rect(s,x+0.03,y+0.018,cellw-0.12,rowh-0.05,fill=fill,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.14)
            text(s,x,y,cellw-0.06,rowh,[[(f"{pct}%  ",{"color":tcol,"size":9.6,"bold":True}),(f"· {n}",{"color":GRAY,"size":8.2})]],
                 align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
            x+=cellw
        y+=rowh
    return y

def circuit(s,t,nodes,cardh=3.35):
    """Circuito horizontal de pasos (sábados). nodes: lista de dicts."""
    l0=0.55; n=len(nodes); gap=0.30; cw=(12.23-gap*(n-1))/n
    for i,nd in enumerate(nodes):
        x=l0+i*(cw+gap); acc=nd["color"]
        rect(s,x,t,cw,cardh,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
        rect(s,x,t,cw,0.6,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.07); rect(s,x,t+0.3,cw,0.3,fill=acc)
        text(s,x+0.1,t,cw-0.2,0.6,nd["tag"],size=12,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
        # puntos + acumulado
        text(s,x+0.2,t+0.72,cw-0.4,0.5,[[(f"{nd['pts']} ",{"color":acc,"size":22,"bold":True}),("pts",{"color":GRAY,"size":11})]])
        text(s,x+0.2,t+1.24,cw-0.4,0.28,f"acumulado {nd['acum']}/108",size=8.6,color=GRAY)
        # barra de progreso acumulado
        bw=cw-0.4
        rect(s,x+0.2,t+1.52,bw,0.09,fill=GRID,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
        rect(s,x+0.2,t+1.52,bw*nd["acum"]/108.0,0.09,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
        # titulo del entregable
        text(s,x+0.2,t+1.72,cw-0.36,0.5,nd["title"],size=10.6,color=INK,bold=True,ls=1.02)
        # referencias
        text(s,x+0.2,t+2.28,cw-0.34,cardh-2.4,bullets(nd["refs"],acc,size=8.5),ls=1.02,sa=3)
        # flecha
        if i<n-1:
            text(s,x+cw-0.04,t+1.15,gap+0.1,0.5,"➜",size=17,color=GRAYL,bold=True,align=PP_ALIGN.CENTER)

# ============================================================ 1 · PORTADA
s=slide()
rect(s,0,0,13.333,7.5,fill=WHITE)
rect(s,0,0,4.9,7.5,fill=NAVY)
rect(s,0,0,4.9,0.11,fill=ORANGE)
rect(s,0.5,0.5,0.2,0.2,fill=BLUE,shape=MSO_SHAPE.DIAMOND)
text(s,0.8,0.42,4.0,0.5,[[("web",{"color":WHITE,"bold":True}),("shooks",{"color":WHITE,"bold":True}),(".",{"color":ORANGE,"bold":True})]],size=26)
text(s,0.55,1.12,4.2,0.35,"Desarrollo de Software · LATAM",size=12.5,color=BLUE)
rect(s,0.57,1.62,2.4,0.04,fill=ORANGE)
text(s,0.55,2.1,4.2,1.6,[[("Propuesta",{"color":WHITE,"bold":True})],[("Comercial",{"color":WHITE,"bold":True})]],size=40,ls=1.0)
text(s,0.55,3.95,4.3,1.4,[[("Screener Financiero AR + US",{"color":GRAYL,"size":15,"bold":True})],[("Base de datos unificada · doble moneda · ratios oficiales — Backoffice Catalaxia",{"color":GRAYL,"size":11.5})]],ls=1.1)
rect(s,0.55,6.05,3.85,0.014,fill=NAVY2)
text(s,0.55,6.2,4.0,0.35,[[("Cliente   ",{"color":GRAYL,"size":11}),("Catalaxia S.A.",{"color":WHITE,"size":11,"bold":True})]])
text(s,0.55,6.6,4.0,0.35,[[("Fecha   ",{"color":GRAYL,"size":11}),("Julio 2026  ·  v4.1",{"color":WHITE,"size":11,"bold":True})]])
# right
text(s,5.35,0.95,7.5,2.1,[[("Iniciá tu",{"color":INK,"bold":True})],[("transformación digital",{"color":INK,"bold":True})],[("con nosotros.",{"color":BLUE,"bold":True})]],size=33,ls=1.04)
rect(s,5.4,3.35,7.1,0.04,fill=BLUE)
text(s,5.35,3.6,7.4,1.0,[[("Propuesta de Solución (v4): ",{"color":INK,"size":13,"bold":True}),("base de datos financiera que unifica Argentina y Estados Unidos, en ARS y USD, con ratios CEDEAR/ADR oficiales y trazabilidad por dato.",{"color":BODY,"size":13})]],ls=1.12)
chips=["Python","Next.js","PostgreSQL","FastAPI","SEC EDGAR","CNV · BYMA","Comafi"]
x=5.4; y=5.12
for i,ch in enumerate(chips):
    w=0.32+len(ch)*0.10
    if x+w>12.95: x=5.4; y+=0.5
    chip(s,x,y,w,ch,accent=[BLUE,ORANGE,GREEN,BLUE_D][i%4]); x+=w+0.16

# ============================================================ 2 · RESUMEN EJECUTIVO
s=slide()
header(s,"Resumen Ejecutivo","Base de datos financiera única en el mercado argentino: 572 empresas en 3 mercados, con doble precio ARS/USD y ratios CEDEAR/ADR oficiales. El pipeline ya está construido, validado y se actualiza solo.")
sy=2.15
for i,(num,lab,acc) in enumerate([("572","Empresas · 3 mercados (S&P 500 + ADR + BYMA)",BLUE),
    ("144","CEDEARs con ratio oficial (Banco Comafi)",ORANGE),
    ("733","Instrumentos con ISIN (modelo de 2 niveles)",GREEN),
    ("1,3%","Dispersión del CCL — valida los ratios de paso",BLUE_D)]):
    stat(s,0.55+i*3.06,sy,2.85,1.5,num,lab,accent=acc)
text(s,0.55,3.95,12.2,0.35,"¿Qué incluye?",size=14,color=INK,bold=True)
col1=["3 universos: 499 S&P 500 (US) + 17 ADR + 56 BYMA-only","Doble precio ARS/USD + CCL calculado desde ADRs","Ratios CEDEAR oficiales (Comafi) — N:1 por instrumento","Ratios/datos ADR: SEC F-6/20-F + JPMorgan (adr.com)"]
col2=["Modelo de 2 niveles: empresa + 733 instrumentos (ISIN)","~16 ratios fundamentales (SEC EDGAR + CNV NIC-29)","Provenance: source_url por dato — auditable","Pipeline replicable (run_all) + jobs (diario/trimestral/anual)"]
text(s,0.7,4.4,6.0,2.4,bullets(col1,BLUE),ls=1.12,sa=6)
text(s,6.9,4.4,6.0,2.4,bullets(col2,GREEN),ls=1.12,sa=6)

# ============================================================ 3 · LO NUEVO DESDE v3
s=slide()
header(s,"Lo nuevo desde la v3","La v3 (junio) era EE.UU. + ADRs desde SEC EDGAR. La v4 suma el lado argentino oficial, la doble moneda y los ratios de conversión — sin perder trazabilidad.",accent=ORANGE)
table(s,0.55,2.05,12.23,["Dimensión","v3 · Junio 2026","v4 · Ahora"],
    [["Universo","553 US + ADR (solo EDGAR)",("572 en 3 mercados: S&P 500 + ADR + BYMA",GREEN_D)],
     ["Precio","USD (EDGAR / yfinance)",("doble ARS/USD + CCL calculado",GREEN_D)],
     ["Ratio CEDEAR","— (no estaba)",("144 oficiales del emisor (Comafi)",GREEN_D)],
     ["Ratio ADR","— (no estaba)",("SEC F-6 / 20-F + JPMorgan",GREEN_D)],
     ["Modelo de datos","1 fila por empresa",("2 niveles: empresa + 733 instrumentos con ISIN",GREEN_D)],
     ["Fuentes AR","— (no estaba)",("CNV NIC-29 · BYMA/IAMC · Comafi",GREEN_D)]],
    cw=[2.4,4.2,5.6],fs=10.5,rh=0.56)
card(s,0.55,5.7,12.23,1.05,"Validación cruzada — el CCL como prueba",
    [[("El CCL (contado con liqui) calculado desde 11 ADRs da una mediana de ~1.578 con solo ",{"size":10.5}),("1,3% de dispersión",{"color":GREEN_D,"bold":True,"size":10.5}),(". Esa coherencia valida, de forma independiente, que los ratios ADR/CEDEAR oficiales están bien cargados.",{"size":10.5})]],
    accent=GREEN,bg=LG,hsize=12.5)

# ============================================================ 4 · RATIOS Y DATOS OFICIALES
s=slide()
header(s,"El diferenciador — datos oficiales con trazabilidad","No es otro screener con scraping frágil. Cada dato tiene su fuente oficial y su source_url; los ratios de conversión vienen del emisor, no “a ojo”.")
table(s,0.55,2.05,12.23,["Dato","Fuente oficial","Cobertura"],
    [["Ratio CEDEAR (N CEDEARs = 1 acción)","Banco Comafi (el emisor)","144 · S&P 500"],
     ["Ratio ADR (1 ADR = N acciones)","SEC — formularios F-6 / 20-F","17 · ADR AR"],
     ["Precio de cierre BYMA","BYMA · IAMC (informe diario oficial)","73 · AR"],
     ["Fundamentales EE.UU.","SEC EDGAR — companyfacts (GAAP+IFRS)","499 · S&P 500"],
     ["Fundamentales BYMA","CNV — estados contables NIC 29","56 · BYMA"],
     ["Dividendos + DR level del ADR","JPMorgan — adr.com (API pública)","17 · ADR AR"]],
    cw=[4.6,4.6,3.0],fs=10.5,rh=0.5)
card(s,0.55,5.55,12.23,1.15,"source_url por dato — auditable ante un cliente o un regulador",
    [[("A diferencia de un Excel armado a mano o de un scraping de portales, cada número del screener guarda de dónde salió. Se puede auditar, y las limitaciones están publicadas (no comparamos BYMA vs S&P 1:1 sin ajuste). Vendemos ",{"size":10.5}),("medido, trazable y honesto",{"color":INK,"bold":True,"size":10.5}),(".",{"size":10.5})]],
    accent=BLUE,bg=LB,hsize=12.5)

# ============================================================ 5 · ALCANCE DEL PROYECTO
s=slide()
header(s,"Alcance del Proyecto")
cols=[("Base de datos & pipeline",BLUE,
    ["572 empresas + 733 instrumentos (modelo 2 niveles)","Doble precio ARS/USD + CCL + precio_fuente","Multi-fuente: SEC · CNV · IAMC · Comafi · JPMorgan","Jobs: diario (precios) · trimestral · mensual · anual","Resiliente: si la CNV cae, difiere y avisa","Migración SQLite → PostgreSQL (esquema definido)"]),
   ("API REST · FastAPI",ORANGE,
    ["GET /empresas (screener, filtros, orden)","GET /empresas/{id} (fundamentales + instrumentos)","GET /instrumentos?tipo=cedear&mercado=BYMA","GET /cobertura (validar_final) — salud del dataset","POST /jobs/{rebuild} · GET /jobs/{id}/status","openapi.json documentado"]),
   ("Frontend · Next.js",GREEN,
    ["Screener: 572 empresas, toggle ARS / USD","Filtros por ratio + por flag de calidad","Detalle: fundamentales + panel de instrumentos","Dividendos del ADR + yield · vista CEDEAR","Comparador (mismo universo) + provenance por dato","Exportar CSV · estados de jobs en vivo"])]
w=3.94
for i,(head,acc,items) in enumerate(cols):
    x=0.55+i*4.06
    rect(s,x,2.0,w,4.7,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
    rect(s,x,2.0,w,0.55,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08)
    rect(s,x,2.28,w,0.27,fill=acc)
    text(s,x+0.15,2.0,w-0.3,0.55,head,size=12.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+0.2,2.7,w-0.36,3.9,bullets(items,acc,size=9.8),ls=1.1,sa=6)

# ============================================================ 6 · STACK TECNOLÓGICO
s=slide()
header(s,"Stack Tecnológico")
def stacktbl(s,x,title,rows,acc):
    rect(s,x,2.0,5.9,0.5,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06)
    text(s,x+0.2,2.0,5.5,0.5,title,size=13,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    y=2.62
    for k,v in rows:
        rect(s,x,y,5.9,0.44,fill=LB3,line=GRID,lw=0.5)
        text(s,x+0.18,y,2.2,0.44,k,size=10,color=GRAY,anchor=MSO_ANCHOR.MIDDLE)
        text(s,x+2.5,y,3.3,0.44,v,size=10,color=INK,bold=True,anchor=MSO_ANCHOR.MIDDLE)
        y+=0.5
stacktbl(s,0.55,"BACKEND",[("Lenguaje","Python 3.11+"),("API","FastAPI + Uvicorn"),("Precios","yfinance + IAMC (BYMA)"),
    ("Fundamentales","SEC EDGAR + CNV NIC-29"),("Ratios oficiales","Comafi · SEC F-6 · JPMorgan"),
    ("Base de datos","PostgreSQL 16 (hoy SQLite)"),("ORM / Scheduler","SQLAlchemy 2.0 · APScheduler")],BLUE)
stacktbl(s,6.9,"FRONTEND",[("Framework","Next.js 14 (App Router)"),("Lenguaje","TypeScript 5"),("Estilos","Tailwind CSS 3"),
    ("Componentes","shadcn/ui"),("Fetch / cache","SWR + React Query"),("Iconos","Tabler Icons"),("Tipografía","Calibri / Inter")],GREEN)
card(s,0.55,6.35,12.23,0.7,"Clave vs. la propuesta original",
    [[("Fuentes oficiales múltiples con trazabilidad (source_url) en lugar de scraping frágil. Riesgo legal y de mantenimiento acotado; el pipeline se reconstruye desde un raw cache sin volver a golpear la red.",{"size":10.5})]],
    accent=ORANGE,bg="FDF0E4",hsize=11.5,head_h=0.32)

# ============================================================ 7 · COBERTURA DEL UNIVERSO
s=slide()
header(s,"Cobertura del Universo","De las fuentes oficiales al dataset listo para producción — 3 mercados, doble moneda, 100 % trazable.")
for i,(num,lab,acc) in enumerate([("572","Empresas con ratios calculados",BLUE),
    ("3","Mercados unificados (US · ADR · BYMA)",ORANGE),
    ("144","CEDEARs con ratio oficial (Comafi)",GREEN),
    ("733","Instrumentos negociables con ISIN",BLUE_D)]):
    stat(s,0.55+i*3.06,2.05,2.85,1.45,num,lab,accent=acc)
text(s,0.55,3.75,12.2,0.35,"Composición del dataset",size=14,color=INK,bold=True)
table(s,0.55,4.15,12.23,["Segmento","Fuente / taxonomía","Estado"],
    [["S&P 500 (large caps EE.UU.)","SEC EDGAR · us-gaap",("✓ 499 cubiertas",GREEN_D)],
     ["ADR Argentina (GGAL, PAM, BMA, YPF, CRESY…)","EDGAR (20-F) + F-6 + JPMorgan",("✓ 17 · ratio ADR oficial",GREEN_D)],
     ["BYMA-only (locales sin ADR)","CNV · NIC 29",("✓ 56 · fundamentales AR",GREEN_D)],
     ["CEDEARs con ratio de conversión","Banco Comafi (emisor)",("✓ 144 · N:1 por instrumento",GREEN_D)],
     ["Precio CEDEAR real de IAMC","BYMA · IAMC",("⚠ hoy teórico (usd × CCL)","B8860B")]],
    cw=[4.6,4.6,3.0],fs=10.2,rh=0.5)

# ============================================================ 8 · [NUEVO] COBERTURA POR RATIO Y FUENTE
s=slide()
header(s,"Cobertura de datos — por ratio y por fuente",
    "No es perfecto: es medido. Para cada ratio mostramos qué % del universo tiene dato en cada fuente (y cuántos instrumentos). Lo que falta se marca con flag + source_url — nunca se completa a ojo.")
heatmap(s,0.55,2.2,12.23,
    [("EDGAR · 499",BLUE),("ADR · 17",ORANGE),("Only-BYMA · 56",GREEN)],
    [("ROE",           [(88,437),(100,17),(93,52)]),
     ("ROA",           [(98,490),(100,17),(100,56)]),
     ("Margen Neto",   [(89,446),(59,10),(93,52)]),
     ("Margen Op.",    [(73,364),(65,11),(89,50)]),
     ("Deuda / EBITDA",[(60,300),(47,8),(68,38)]),
     ("EPS",           [(98,487),(88,15),(89,50)]),
     ("FCF / CapEmpl.",[(84,419),(65,11),(91,51)]),
     ("Payout",        [(73,362),(47,8),(61,34)]),
     ("CAGR EPS 5y",   [(69,345),(12,2),(61,34)]),
     ("PER",           [(93,464),(65,11),(34,19)]),
     ("P / Book",      [(87,435),(88,15),(64,36)]),
     ("P / Sales",     [(91,452),(53,9),(64,36)]),
     ("EV / EBITDA",   [(72,357),(47,8),(43,24)])],
    rowh=0.285)
# leyenda
lg=[("≥ 85 % completo",GREEN_D,LG),("70–84 % alto",BLUE_D,LB),("50–69 % parcial",HM_AMBER_T,HM_AMBER),("< 50 % limitado",HM_RED_T,HM_RED)]
lx=0.6
for lab,tc,fc in lg:
    rect(s,lx,6.58,0.22,0.22,fill=fc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.3)
    text(s,lx+0.3,6.55,2.6,0.3,lab,size=9,color=GRAY,anchor=MSO_ANCHOR.MIDDLE); lx+=2.9
text(s,0.6,6.86,12.2,0.3,[[("Precio y market cap: 100 % / 88 % / 96 %.  ",{"color":INK,"size":9,"bold":True}),
    ("Los múltiplos (PER, EV/EBITDA) bajan en BYMA porque dependen del precio local; el CAGR del ADR es el punto más flaco (falta histórico). Cobertura media global ≈ 84 %.",{"color":GRAY,"size":9})]],ls=1.0)

# ============================================================ 9 · ARQUITECTURA DEL PIPELINE
s=slide()
header(s,"Arquitectura del Pipeline","Cada bloque tiene input/output definido y se re-ejecuta solo. El raw cache de cada fuente se descarga una vez (lo caro es el rate limit, no el cómputo).")
blocks=[("1 · Extracción",BLUE,"SEC EDGAR · CNV · IAMC · Comafi · JPMorgan → raw cache local (fuente de verdad)"),
    ("2 · Ratios base",GREEN,"~16 ratios fundamentales: márgenes, ROE, deuda/EBITDA, EV/EBITDA, CAGR"),
    ("3 · Ensamble AR",ORANGE,"doble precio ARS/USD, CCL, ratios CEDEAR (Comafi) y ADR (F-6/JPM)"),
    ("4 · Modelo 2 niveles",BLUE_D,"empresa + 733 instrumentos con ISIN, moneda, precio y ratio"),
    ("5 · Calidad + Export",NAVY2,"flags / sanity gates / no_significativo → CSV + JSON + SQL con source_url")]
w=2.32
for i,(t,acc,d) in enumerate(blocks):
    x=0.55+i*2.46
    rect(s,x,2.15,w,2.5,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
    rect(s,x,2.15,w,0.62,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06)
    rect(s,x,2.46,w,0.31,fill=acc)
    text(s,x+0.1,2.15,w-0.2,0.62,t,size=11.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
    text(s,x+0.15,2.9,w-0.3,1.6,d,size=9.4,color=BODY,ls=1.08)
    if i<4:
        text(s,x+w-0.02,3.1,0.5,0.4,"→",size=18,color=GRAYL,bold=True)
card(s,0.55,5.05,12.23,1.65,"Principio de diseño + resiliencia",
    [[("El raw cache es la fuente de verdad: se baja una vez de cada fuente y la base se reconstruye sin volver a la red. Si cambia una fórmula, se recalcula en segundos.",{"size":10.5})],
     [("Resiliencia CNV (single point of failure): si la fuente argentina no responde, el pipeline ",{"size":10.5}),("difiere y avisa",{"color":ORANGE,"bold":True,"size":10.5}),(" (cnv_last_status) en vez de arrastrar dato viejo en silencio. Jobs periódicos: precios diario · fundamentales trimestral · IPC mensual · universo anual.",{"size":10.5})]],
    accent=ORANGE,bg="FDF0E4",hsize=12.5,head_h=0.34,bls=1.12)

# ============================================================ 10 · EQUIPO
s=slide()
header(s,"Equipo de Trabajo")
team=[("Federico — Arquitecto · Tech Lead",BLUE_D,["Arquitectura del pipeline multi-fuente","Modelo de 2 niveles (empresa + instrumentos)","Metodología de validación (CCL, provenance)","Integración, code review y QA","Deploy y soporte post-entrega"]),
    ("Mateo — Backend",BLUE,["Pipeline SEC EDGAR + CNV NIC-29","Ratios CEDEAR/ADR oficiales (Comafi, F-6, JPM)","Doble precio ARS/USD + CCL","Jobs periódicos + resiliencia CNV","Endpoints FastAPI + APScheduler"]),
    ("Joaquín — Frontend A",GREEN,["Screener: tabla 572 empresas, toggle ARS/USD","Filtros por ratio + flag de calidad","Sorting + paginación + exportar CSV","Panel de jobs en tiempo real","Detalle de empresa"]),
    ("Aldana — Frontend B",ORANGE,["Panel de instrumentos (ADR/CEDEAR/local)","Dividendos del ADR + yield","Vista CEDEAR + comparador","Chips de provenance por dato","Estados vacíos / loading / responsive"])]
w=5.95
for i,(name,acc,items) in enumerate(team):
    x=0.55+(i%2)*6.28; y=2.0+(i//2)*2.5
    rect(s,x,y,w,2.32,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
    rect(s,x,y,0.08,2.32,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,x+0.22,y+0.12,w-0.4,0.4,name,size=12.5,color=acc,bold=True)
    text(s,x+0.28,y+0.58,w-0.5,1.65,bullets(items,acc,size=9.4),ls=1.05,sa=3)

# ============================================================ 11 · CRONOGRAMA
s=slide()
header(s,"Cronograma — 4 semanas de productización","Fase 0 — la data ya está construida y validada (572 empresas, 3 mercados, ratios oficiales, CCL 1,3 %). El riesgo técnico más grande ya está resuelto: solo falta empaquetarla.",accent=GREEN)
weeks=[("Fase 0 · Data",GREEN,["✓ 572 empresas, 3 mercados","✓ Ratios CEDEAR/ADR oficiales","✓ Doble precio + CCL validado","✓ Modelo 2 niveles + provenance"]),
    ("Semana 1 · Fundamentos",BLUE,["Migración esquema a PostgreSQL","Endpoints base + jobs","Next.js + sidebar/topbar"]),
    ("Semana 2 · Datos",BLUE,["Servir 572 empresas vía API","Screener con filtros reales","Toggle ARS/USD + detalle"]),
    ("Semana 3 · Completo",ORANGE,["Panel de instrumentos + dividendos","Vista CEDEAR + comparador","Jobs de actualización automática"]),
    ("Semana 4 · QA",NAVY2,["Tests + manejo de errores","openapi.json documentado","QA y entrega final"])]
w=2.32
for i,(t,acc,items) in enumerate(weeks):
    x=0.55+i*2.46
    rect(s,x,2.15,w,3.9,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
    rect(s,x,2.15,w,0.6,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06)
    rect(s,x,2.45,w,0.3,fill=acc)
    text(s,x+0.1,2.15,w-0.2,0.6,t,size=10.8,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
    text(s,x+0.16,2.9,w-0.3,3.0,bullets(items,acc,size=9.2),ls=1.06,sa=5)
card(s,0.55,6.2,12.23,0.62,"Entregas los sábados",[[("Cada sábado: rama, demo y checklist con referencias (épica · sprint · story points). Detalle en las 2 láminas siguientes.",{"size":10.5})]],accent=BLUE,bg=LB,hsize=11.5,head_h=0.3)

# ============================================================ 12 · [NUEVO] CIRCUITO DE ENTREGAS POR SÁBADO
s=slide()
header(s,"Entregas de los sábados — un sábado, un avance",
    "108 puntos de historia en 34 tareas. Cada sábado cerramos un bloque navegable y lo entregamos con sus referencias marcadas. El acumulado sube 38 → 63 → 89 → 108.",accent=GREEN)
circuit(s,2.15,[
    {"tag":"Sábado 1 · Sem 1","pts":38,"acum":38,"color":BLUE,
     "title":"Fundamentos navegables","refs":[("Infra","Postgres + 572 empresas migradas"),("API base","/empresas · /cobertura · /health + HTTPS"),("Front","screener + detalle contra mocks")]},
    {"tag":"Sábado 2 · Sem 2","pts":25,"acum":63,"color":GREEN,
     "title":"Datos reales conectados","refs":[("Backend","/empresas/{id} + /instrumentos"),("Screener","API real + toggle ARS/USD"),("Detalle","panel de instrumentos")]},
    {"tag":"Sábado 3 · Sem 3","pts":26,"acum":89,"color":ORANGE,
     "title":"Funcionalidad completa","refs":[("Filtros","compuestos + export CSV"),("Jobs","scheduler + panel en vivo"),("CEDEAR/ADR","dividendos + comparador")]},
    {"tag":"Sábado 4 · Sem 4","pts":19,"acum":108,"color":BLUE_D,
     "title":"QA + entrega","refs":[("QA","tests + e2e + openapi.json"),("Infra","hardening + backups + prod"),("Handover","walkthrough + README")]},
])
card(s,0.55,5.75,12.23,1.0,"Referencias bien marcadas",
    [[("Cada entregable del sábado cita su ",{"size":10.3}),("épica · sprint · story points",{"color":INK,"bold":True,"size":10.3}),
      (" del backlog (34 tareas). El cliente sigue el avance en Jira y en la demo — sin sorpresas al final, y con la trazabilidad de siempre.",{"size":10.3})]],
    accent=GREEN,bg=LG,hsize=12,head_h=0.32)

# ============================================================ 13 · [NUEVO] LOS 108 PUNTOS POR ÉPICA
s=slide()
header(s,"Los 108 puntos — desglose por épica",
    "34 tareas · 7 épicas · 4 sprints. Así se referencia cada entrega de los sábados.")
epics=[("Backend · API REST","Mateo",8,27,BLUE),
       ("Frontend · Screener","Joaquín",8,24,GREEN),
       ("Frontend · Detalle","Aldana",7,22,ORANGE),
       ("Infra & Deploy (VPS)","Federico",4,13,BLUE_D),
       ("Datos & Base de datos","Federico",3,11,BLUE),
       ("Jobs & Scheduler","Mateo · Fede",2,6,GREEN),
       ("QA & Entrega","Federico",2,5,ORANGE)]
# encabezado tabla
tl,tt,tw=0.55,2.0,12.23
heads=["Épica","Responsable","Tareas","Puntos","Peso (proporcional)"]
cw=[3.5,2.3,1.2,1.2,4.03]; hrh=0.44; rh=0.44
x=tl
for j,hd in enumerate(heads):
    rect(s,x,tt,cw[j],hrh,fill=NAVY)
    text(s,x+0.1,tt,cw[j]-0.14,hrh,hd,size=10,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    x+=cw[j]
y=tt+hrh
maxp=max(e[3] for e in epics)
for i,(ep,resp,nt,pts,acc) in enumerate(epics):
    rect(s,tl,y,tw,rh,fill=(LB3 if i%2==0 else WHITE),line=GRID,lw=0.5)
    x=tl
    text(s,x+0.1,y,cw[0]-0.14,rh,[[("● ",{"color":acc,"size":10}),(ep,{"color":INK,"size":10,"bold":True})]],anchor=MSO_ANCHOR.MIDDLE); x+=cw[0]
    text(s,x+0.1,y,cw[1]-0.14,rh,resp,size=9.6,color=BODY,anchor=MSO_ANCHOR.MIDDLE); x+=cw[1]
    text(s,x+0.1,y,cw[2]-0.14,rh,str(nt),size=9.6,color=BODY,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER); x+=cw[2]
    text(s,x+0.1,y,cw[3]-0.14,rh,[[(str(pts),{"color":acc,"size":11,"bold":True})]],anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER); x+=cw[3]
    bw=(cw[4]-0.35)*pts/maxp
    rect(s,x+0.12,y+rh/2-0.11,cw[4]-0.35,0.22,fill=GRID,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    rect(s,x+0.12,y+rh/2-0.11,bw,0.22,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    y+=rh
# fila total
rect(s,tl,y,tw,rh,fill=NAVY)
text(s,tl+0.1,y,cw[0]-0.14,rh,"TOTAL",size=10.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE)
text(s,tl+cw[0]+0.1,y,cw[1]-0.14,rh,"7 épicas",size=9.6,color=GRAYL,anchor=MSO_ANCHOR.MIDDLE)
text(s,tl+cw[0]+cw[1]+0.1,y,cw[2]-0.14,rh,"34",size=10.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
text(s,tl+cw[0]+cw[1]+cw[2]+0.1,y,cw[3]-0.14,rh,"108",size=12,color=ORANGE,bold=True,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
# distribución por sprint / sábado
sy=y+rh+0.30
text(s,0.55,sy-0.28,8,0.3,"Distribución por sprint (un sábado por sprint)",size=11,color=INK,bold=True)
sp=[("Sáb 1 · Sprint 1",38,BLUE),("Sáb 2 · Sprint 2",25,GREEN),("Sáb 3 · Sprint 3",26,ORANGE),("Sáb 4 · Sprint 4",19,BLUE_D)]
sw=(12.23-0.3*3)/4
for i,(lab,pts,acc) in enumerate(sp):
    x=0.55+i*(sw+0.3)
    rect(s,x,sy+0.02,sw,0.6,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08)
    rect(s,x,sy+0.02,0.07,0.6,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,x+0.2,sy+0.08,sw-0.3,0.26,lab,size=9.3,color=GRAY)
    text(s,x+0.2,sy+0.3,sw-0.3,0.3,[[(str(pts),{"color":acc,"size":14,"bold":True}),(" pts",{"color":GRAY,"size":9.5})]])

# ============================================================ 14 · INVERSIÓN
s=slide()
header(s,"Inversión y Financiación")
rect(s,0.55,2.05,5.7,3.1,fill=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
text(s,0.85,2.35,5.1,0.4,"Inversión total",size=13,color=GRAYL)
text(s,0.85,2.75,5.1,1.0,"USD 1.800",size=44,color=WHITE,bold=True)
text(s,0.85,3.8,5.1,0.4,"Precio fijo por proyecto · sin costos ocultos",size=11,color=BLUE)
text(s,0.85,4.25,5.1,0.8,[[("Incluye la data ya construida y validada — 572 empresas, 3 mercados, ratios oficiales — sin cargo extra.",{"color":GRAYL,"size":10.5})]],ls=1.1)
text(s,0.85,4.95,5.1,0.35,"Soporte correctivo: USD 20/mes o pago único USD 200",size=9.5,color=GRAYL,italic=True)
text(s,6.5,2.05,6.3,0.35,"Opciones de pago",size=14,color=INK,bold=True)
pays=[("Contado",BLUE,"USD 1.800","Pago único al firmar · 10 % off si se contrata hoy"),
    ("3 cuotas sin interés",GREEN,"USD 600 × 3","Al firmar · a la entrega · mes posterior"),
    ("6 cuotas",ORANGE,"USD 320 × 6","Al firmar · a la entrega · mensuales")]
y=2.5
for t,acc,amt,det in pays:
    rect(s,6.5,y,6.3,0.82,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06)
    rect(s,6.5,y,0.08,0.82,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,6.7,y+0.12,3.0,0.3,t,size=12,color=acc,bold=True)
    text(s,6.7,y+0.44,4.2,0.3,det,size=9,color=GRAY)
    text(s,10.2,y,2.5,0.82,amt,size=16,color=INK,bold=True,align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
    y+=0.95
text(s,6.5,5.5,6.3,0.5,"Precio orientativo por el alcance ampliado (3 mercados, doble moneda, ratios oficiales); ajustable según el alcance final acordado.",size=9,color=GRAY,italic=True,ls=1.05)

# ============================================================ 15 · CÓMO ARRANCAMOS
s=slide()
header(s,"¿Cómo arrancamos?")
steps=[("01","Aprobación","Confirmar la propuesta y el plan de pago.",BLUE),
    ("02","Primer pago","50 % al inicio o primera cuota según el plan.",GREEN),
    ("03","Kickoff","1 h con el equipo: recorrido por las 572 empresas ya validadas + contrato de API.",ORANGE),
    ("04","Entrega","4 semanas de desarrollo + entregas cada sábado.",BLUE_D)]
w=2.95
for i,(n,t,d,acc) in enumerate(steps):
    x=0.55+i*3.07
    rect(s,x,2.1,w,2.0,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
    text(s,x+0.2,2.25,1.5,0.7,n,size=30,color=acc,bold=True)
    text(s,x+0.2,2.95,w-0.4,0.35,t,size=13,color=INK,bold=True)
    text(s,x+0.2,3.35,w-0.4,0.7,d,size=9.6,color=BODY,ls=1.08)
# resumen box
rect(s,0.55,4.45,12.23,1.55,fill=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.03)
items=[("Empresa","webshooks"),("Inversión","USD 1.800"),("Proyecto","Screener AR + US · doble moneda · ratios oficiales"),
    ("Duración","4 semanas"),("Inicio","Inmediato"),("Garantía","30 días post-entrega")]
for i,(k,v) in enumerate(items):
    x=0.9+(i%3)*4.05; y=4.7+(i//3)*0.62
    text(s,x,y,3.9,0.3,[[(k+"   ",{"color":GRAYL,"size":10}),(v,{"color":WHITE,"size":11,"bold":True})]])
text(s,0.55,6.35,12.2,0.4,"Esta propuesta tiene validez de 7 días.",size=10.5,color=GRAY,italic=True,align=PP_ALIGN.CENTER)

import os
outs=[r"D:/Software Development/Porfolio/catalaxia-cedears-prod/docs/pdfs/Propuesta_webshooks_Catalaxia_v4.1.pptx",
      r"D:/Software Development/Porfolio/catalaxia-cedears-prod/docs/proyecto/Propuesta_webshooks_Catalaxia_v4.1.pptx"]
for out in outs:
    try:
        prs.save(out); print("SAVED:",out,"| slides:",len(prs.slides._sldIdLst))
    except Exception as e:
        print("NO GUARDADO (¿abierto en PowerPoint?):",out,"->",e)
