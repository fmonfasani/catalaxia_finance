# -*- coding: utf-8 -*-
"""Kickoff tecnico del equipo Catalaxia -> PPTX branded webshooks.
Que construimos, arquitectura, contrato de API, modelo de datos, como trabajamos,
roles y las tareas de cada uno por sprint. Mismo estilo que la propuesta v4."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY="1A2744"; NAVY2="3A4664"; BLUE="3B82F6"; BLUE_D="1A4D8F"
ORANGE="F5821F"; GREEN="43A047"; GREEN_D="1F6B35"
LB="DCEAFF"; LG="E7F7EC"; LB3="F4F8FE"; GRID="D6E2F2"
BODY="3A4458"; INK="1A2744"; GRAY="6B7589"; GRAYL="AAB4C8"; WHITE="FFFFFF"
FONT="Calibri"
ACC={"Federico":BLUE_D,"Mateo":BLUE,"Joaquín":GREEN,"Aldana":ORANGE}

def C(h): return RGBColor.from_string(h)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]; PG={"n":1}
def slide(): return prs.slides.add_slide(BLANK)
def _ns(sp):
    try: sp.shadow.inherit=False
    except: pass
def rect(s,l,t,w,h,fill=None,line=None,lw=1.0,shape=MSO_SHAPE.RECTANGLE,rad=None):
    sp=s.shapes.add_shape(shape,Inches(l),Inches(t),Inches(w),Inches(h)); _ns(sp)
    if fill is None: sp.fill.background()
    else: sp.fill.solid(); sp.fill.fore_color.rgb=C(fill)
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=C(line); sp.line.width=Pt(lw)
    if rad is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0]=rad
        except: pass
    return sp
def text(s,l,t,w,h,runs,size=11,color=BODY,bold=False,italic=False,align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP,ls=1.0,sa=2):
    tb=s.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); tf=tb.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,Pt(1))
    if isinstance(runs,str): runs=[runs]
    for i,para in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.line_spacing=ls; p.space_after=Pt(sa); p.space_before=Pt(0)
        if isinstance(para,str): para=[(para,{})]
        for seg in para:
            txt,kw=seg; r=p.add_run(); r.text=txt
            r.font.size=Pt(kw.get("size",size)); r.font.name=FONT
            r.font.color.rgb=C(kw.get("color",color)); r.font.bold=kw.get("bold",bold)
            r.font.italic=kw.get("italic",italic)
    return tb
def header(s,title,subtitle=None,accent=BLUE):
    rect(s,0,0,0.09,7.5,fill=accent)
    text(s,0.55,0.30,4,0.3,[[("web",{"color":BLUE,"bold":True}),("shooks",{"color":INK,"bold":True}),(".",{"color":ORANGE,"bold":True})]],size=15)
    text(s,0.55,0.72,11.5,0.6,title,size=25,color=INK,bold=True)
    rect(s,0.57,1.32,0.9,0.045,fill=ORANGE)
    if subtitle: text(s,0.57,1.42,12.2,0.55,subtitle,size=11.5,color=GRAY,ls=1.02)
    rect(s,0.55,7.06,12.23,0.014,fill=GRID)
    text(s,0.55,7.1,4,0.3,"webshooks.com · Kickoff Catalaxia",size=9,color=GRAYL)
    text(s,12.2,7.1,0.6,0.3,str(PG["n"]),size=9,color=GRAYL,align=PP_ALIGN.RIGHT); PG["n"]+=1
def card(s,l,t,w,h,head,body,accent=BLUE,bg=LB3,hsize=12,bsize=9.6,head_h=0.4,bls=1.05):
    rect(s,l,t,w,h,fill=bg,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
    rect(s,l,t,0.07,h,fill=accent,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,l+0.2,t+0.1,w-0.35,head_h,head,size=hsize,color=accent,bold=True)
    if body is not None: text(s,l+0.2,t+0.1+head_h,w-0.38,h-head_h-0.2,body,size=bsize,color=BODY,ls=bls,sa=3)
def stat(s,l,t,w,h,num,lab,accent=BLUE):
    rect(s,l,t,w,h,fill=LB,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08)
    rect(s,l,t,w,0.06,fill=accent)
    text(s,l,t+0.16,w,0.7,num,size=30,color=INK,bold=True,align=PP_ALIGN.CENTER)
    text(s,l+0.1,t+0.92,w-0.2,h-0.98,lab,size=10,color=GRAY,align=PP_ALIGN.CENTER,ls=0.98)
def chip(s,l,t,w,txt,accent=BLUE):
    rect(s,l,t,w,0.36,fill=LB,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,l,t+0.05,w,0.28,txt,size=9.5,color=BLUE_D,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
def table(s,l,t,w,headers,rows,cw=None,fs=9.5,rh=0.44,hrh=0.44,accent=NAVY):
    n=len(headers); cw=[w/n]*n if not cw else [c*(w/sum(cw)) for c in cw]; x=l
    for j,hd in enumerate(headers):
        rect(s,x,t,cw[j],hrh,fill=accent); text(s,x+0.08,t,cw[j]-0.12,hrh,hd,size=fs+0.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,ls=0.9); x+=cw[j]
    y=t+hrh
    for i,row in enumerate(rows):
        rect(s,l,y,w,rh,fill=(LB3 if i%2==0 else WHITE),line=GRID,lw=0.5); x=l
        for j,cell in enumerate(row):
            col=INK if j==0 else BODY; bold=(j==0)
            if isinstance(cell,tuple): cell,col=cell; bold=True
            text(s,x+0.08,y,cw[j]-0.12,rh,str(cell),size=fs,color=col,bold=bold,anchor=MSO_ANCHOR.MIDDLE,ls=0.92); x+=cw[j]
        y+=rh
    return y
def bullets(items,accent=BLUE,size=10.2):
    out=[]
    for it in items:
        if isinstance(it,tuple):
            out.append([("▸  ",{"color":accent,"bold":True,"size":size}),(it[0]+"  ",{"color":INK,"bold":True,"size":size}),(it[1],{"color":BODY,"size":size})])
        else:
            out.append([("▸  ",{"color":accent,"bold":True,"size":size}),(it,{"color":BODY,"size":size})])
    return out

# backlog (sprint, tarea, epica, resp, pts)
B=[
 ("S1","Definir contrato de API + mocks (día 1)","Backend","Federico",3),
 ("S1","Esquema PostgreSQL (empresas + instrumentos)","Datos","Federico",3),
 ("S1","Migrar 572 empresas + 733 instrumentos a Postgres","Datos","Federico",5),
 ("S1","Postgres en la VPS + backups diarios","Infra","Federico",3),
 ("S1","Deploy API base + HTTPS + CORS","Infra","Federico",5),
 ("S1","Scaffold FastAPI + Postgres async","Backend","Mateo",3),
 ("S1","GET /empresas + /cobertura + /health","Backend","Mateo",5),
 ("S1","Scaffold Next.js + layout","Screener","Joaquín",3),
 ("S1","Tabla screener contra mocks","Screener","Joaquín",3),
 ("S1","Base del front + ruta detalle (mocks)","Detalle","Aldana",3),
 ("S1","Componentes base (cards, chips, tabla)","Detalle","Aldana",2),
 ("S2","GET /empresas/{id} (fund. + instrumentos)","Backend","Mateo",5),
 ("S2","GET /instrumentos con filtros + paginación","Backend","Mateo",3),
 ("S2","Verificar 733 instrumentos + dividendos ADR","Datos","Federico",3),
 ("S2","Asistir a frontends (CORS/staging)","Infra","Federico",2),
 ("S2","Screener contra API real + toggle ARS/USD","Screener","Joaquín",5),
 ("S2","Sorting + paginación server-side","Screener","Joaquín",2),
 ("S2","Detalle real + panel de instrumentos","Detalle","Aldana",5),
 ("S3","Filtros compuestos en /empresas","Backend","Mateo",3),
 ("S3","POST /jobs/{tipo} + GET /jobs/{id}/status","Jobs","Mateo",3),
 ("S3","Scheduler de jobs en la VPS","Jobs","Federico",3),
 ("S3","Filtros min/max + flag de calidad","Screener","Joaquín",3),
 ("S3","Export CSV del screener","Screener","Joaquín",2),
 ("S3","Panel de jobs en vivo (polling)","Screener","Joaquín",3),
 ("S3","Dividendos ADR + yield","Detalle","Aldana",3),
 ("S3","Vista CEDEAR (ratio de conversión)","Detalle","Aldana",3),
 ("S3","Comparador + chips de provenance","Detalle","Aldana",3),
 ("S4","Tests de endpoints + manejo de errores","Backend","Mateo",3),
 ("S4","openapi.json documentado","Backend","Mateo",2),
 ("S4","Hardening infra + backups + deploy prod","Infra","Federico",3),
 ("S4","QA de integración end-to-end","QA","Federico",3),
 ("S4","Estados + responsive (screener)","Screener","Joaquín",3),
 ("S4","Estados + responsive + disclaimers (detalle)","Detalle","Aldana",3),
 ("S4","Walkthrough + README de operación","QA","Federico",2),
]
PEOPLE=["Federico","Mateo","Joaquín","Aldana"]

# ============================================================ 1 PORTADA
s=slide(); rect(s,0,0,13.333,7.5,fill=WHITE); rect(s,0,0,4.9,7.5,fill=NAVY); rect(s,0,0,4.9,0.11,fill=ORANGE)
rect(s,0.5,0.5,0.2,0.2,fill=BLUE,shape=MSO_SHAPE.DIAMOND)
text(s,0.8,0.42,4,0.5,[[("web",{"color":WHITE,"bold":True}),("shooks",{"color":WHITE,"bold":True}),(".",{"color":ORANGE,"bold":True})]],size=26)
text(s,0.55,1.12,4.2,0.35,"Desarrollo de Software · LATAM",size=12.5,color=BLUE)
rect(s,0.57,1.62,2.4,0.04,fill=ORANGE)
text(s,0.55,2.1,4.3,1.6,[[("Kickoff",{"color":WHITE,"bold":True})],[("Técnico",{"color":WHITE,"bold":True})]],size=40,ls=1.0)
text(s,0.55,3.95,4.3,1.4,[[("Screener Financiero AR + US",{"color":GRAYL,"size":15,"bold":True})],[("Plan de trabajo · 4 sprints · equipo webshooks",{"color":GRAYL,"size":11.5})]],ls=1.1)
rect(s,0.55,6.05,3.85,0.014,fill=NAVY2)
text(s,0.55,6.2,4,0.35,[[("Duración   ",{"color":GRAYL,"size":11}),("4 semanas",{"color":WHITE,"size":11,"bold":True})]])
text(s,0.55,6.6,4,0.35,[[("Fecha   ",{"color":GRAYL,"size":11}),("Julio 2026",{"color":WHITE,"size":11,"bold":True})]])
text(s,5.35,0.95,7.5,2.1,[[("Construimos una",{"color":INK,"bold":True})],[("plataforma de análisis",{"color":INK,"bold":True})],[("fundamental en 4 semanas.",{"color":BLUE,"bold":True})]],size=31,ls=1.05)
rect(s,5.4,3.35,7.1,0.04,fill=BLUE)
text(s,5.35,3.6,7.4,1.0,[[("Objetivo:  ",{"color":INK,"size":13,"bold":True}),("productizar un dataset ya construido y validado (572 empresas, 3 mercados, ratios oficiales) en un screener web con API propia.",{"color":BODY,"size":13})]],ls=1.12)
x=5.4; y=5.12
for i,ch in enumerate(["FastAPI","Next.js","PostgreSQL","Docker","Jira","Contrato-first"]):
    w=0.32+len(ch)*0.10
    if x+w>12.95: x=5.4; y+=0.5
    chip(s,x,y,w,ch,accent=[BLUE,ORANGE,GREEN,BLUE_D][i%4]); x+=w+0.16

# ============================================================ 2 QUE CONSTRUIMOS
s=slide(); header(s,"Qué vamos a construir","Un screener de análisis fundamental que unifica Argentina y Estados Unidos, en ARS y USD, sobre datos oficiales y trazables.")
for i,(n,l,a) in enumerate([("572","empresas · 3 mercados",BLUE),("733","instrumentos con ISIN",ORANGE),("~16","ratios fundamentales",GREEN),("2","niveles: empresa + instrumento",BLUE_D)]):
    stat(s,0.55+i*3.06,2.15,2.85,1.5,n,l,accent=a)
text(s,0.55,3.95,12,0.35,"Dos vistas de producto",size=14,color=INK,bold=True)
text(s,0.7,4.4,6,2.4,bullets([("Screener","tabla filtrable de 572 empresas, toggle ARS/USD, orden por ratio."),("Ficha de empresa","fundamentales + panel de instrumentos (ADR/CEDEAR/local)."),("Export","CSV del screener.")],BLUE),ls=1.14,sa=7)
text(s,6.9,4.4,6,2.4,bullets([("Datos oficiales","SEC · CNV · IAMC · Comafi · JPMorgan, con source_url por dato."),("Doble moneda","precio ARS/USD + CCL; ratios CEDEAR/ADR del emisor."),("Provenance","cada número es auditable — el diferenciador.")],GREEN),ls=1.14,sa=7)

# ============================================================ 3 PUNTO DE PARTIDA
s=slide(); header(s,"Punto de partida","La capa de datos ya está construida y validada. El trabajo de estas 4 semanas es productizarla: PostgreSQL + API + frontend.",accent=GREEN)
text(s,0.7,2.15,12,2.4,bullets([
    ("Dataset listo","572 empresas + 733 instrumentos, ~16 ratios calculados."),
    ("Ratios oficiales","CEDEAR (Comafi) y ADR (SEC F-6/20-F + JPMorgan)."),
    ("Trazabilidad","source_url por dato; validación cruzada por CCL (dispersión 1,3%)."),
    ("Hoy en SQLite","se migra a PostgreSQL en la Semana 1."),
],GREEN,size=11.5),ls=1.2,sa=10)
card(s,0.55,5.35,12.23,1.35,"Qué implica para el equipo",
    [[("El riesgo de datos ya está resuelto. El foco de las 4 semanas es ingeniería de producto: esquema Postgres, API REST, dos frontends y jobs de actualización. El ",{"size":10.5}),("contrato de API del día 1",{"color":INK,"bold":True,"size":10.5}),(" permite que backend y frontend avancen en paralelo desde el inicio.",{"size":10.5})]],
    accent=BLUE,bg=LB,hsize=12.5)

# ============================================================ 4 ARQUITECTURA
s=slide(); header(s,"Arquitectura objetivo","Cada bloque tiene dueño y contrato. Los frontends consumen la API; la API sirve desde Postgres; los jobs actualizan solos.")
blocks=[("Datos · PostgreSQL",BLUE_D,"Federico","572/733 + provenance;\nmigración SQLite→PG; backups"),
        ("API · FastAPI",BLUE,"Mateo","REST + filtros + jobs;\nOpenAPI documentada"),
        ("Front · Screener",GREEN,"Joaquín","tabla, filtros, toggle\nARS/USD, export, jobs"),
        ("Front · Detalle",ORANGE,"Aldana","ficha, instrumentos,\nCEDEAR/ADR, comparador")]
w=2.9
for i,(t,acc,who,d) in enumerate(blocks):
    x=0.55+i*(w+0.19)
    rect(s,x,2.15,w,2.7,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
    rect(s,x,2.15,w,0.62,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06); rect(s,x,2.46,w,0.31,fill=acc)
    text(s,x+0.12,2.15,w-0.24,0.62,t,size=11.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE,align=PP_ALIGN.CENTER)
    text(s,x+0.15,2.9,w-0.3,0.3,who,size=11,color=acc,bold=True)
    text(s,x+0.15,3.3,w-0.3,1.4,d,size=9.6,color=BODY,ls=1.12)
    if i<3: text(s,x+w-0.02,3.2,0.4,0.4,"→",size=18,color=GRAYL,bold=True)
card(s,0.55,5.2,12.23,1.5,"Infraestructura (Federico)",
    [[("Docker Compose · VPS · HTTPS (nginx/certbot) · CORS · scheduler de jobs · CI (lint + tests en cada PR). ",{"size":10.5}),("El raw cache es la fuente de verdad: la base se reconstruye sin volver a la red.",{"color":INK,"bold":True,"size":10.5})]],
    accent=ORANGE,bg="FDF0E4",hsize=12.5,head_h=0.34)

# ============================================================ 5 CONTRATO API
s=slide(); header(s,"Contrato de API primero (día 1)","Definimos el contrato y entregamos mocks el día 1. Frontend y backend trabajan en paralelo desde el minuto cero — nadie espera a nadie.")
table(s,0.55,2.05,12.23,["Endpoint","Qué devuelve","Dueño"],
    [["GET /empresas","screener: lista + filtros + orden","Mateo"],
     ["GET /empresas/{id}","fundamentales + instrumentos anidados","Mateo"],
     ["GET /instrumentos","filtros por tipo/mercado + paginación","Mateo"],
     ["GET /cobertura","salud del dataset (validar_final)","Mateo"],
     ["POST /jobs/{tipo}  ·  GET /jobs/{id}/status","disparar y seguir jobs","Mateo"],
     ["GET /health","liveness","Mateo"]],
    cw=[4.6,5.6,2.0],fs=10.5,rh=0.52)
card(s,0.55,5.75,12.23,0.95,"Regla de oro",
    [[("El contrato + los mocks se congelan el día 1. Cualquier cambio se acuerda entre backend y frontend y se versiona. ",{"size":10.5}),("Los frontends nunca quedan bloqueados esperando la API.",{"color":INK,"bold":True,"size":10.5})]],
    accent=BLUE,bg=LB,hsize=12,head_h=0.32)

# ============================================================ 6 MODELO DE DATOS
s=slide(); header(s,"Modelo de datos — 2 niveles","Una empresa tiene N instrumentos negociables. Los fundamentales viven en la empresa; el precio y el ratio de conversión, en el instrumento.")
def minitbl(x,title,rows,acc):
    rect(s,x,2.05,5.9,0.5,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.06)
    text(s,x+0.2,2.05,5.5,0.5,title,size=13,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE); y=2.67
    for k,v in rows:
        rect(s,x,y,5.9,0.46,fill=LB3,line=GRID,lw=0.5)
        text(s,x+0.18,y,2.3,0.46,k,size=10,color=GRAY,anchor=MSO_ANCHOR.MIDDLE)
        text(s,x+2.6,y,3.2,0.46,v,size=10,color=INK,bold=True,anchor=MSO_ANCHOR.MIDDLE); y+=0.52
minitbl(0.55,"EMPRESA (nivel 1)",[("id / cuit / cik","identidad"),("nombre · sector · país","taxonomía"),("~16 ratios","fundamentales"),("provenance","source_url por dato"),("fuente_fund","edgar / cnv")],BLUE)
minitbl(6.9,"INSTRUMENTO (nivel 2)",[("tipo","ADR / CEDEAR / local"),("ISIN · mercado · moneda","identidad"),("ratio","conversión (N:1)"),("precio","ARS y USD"),("dividendos / yield","ADR")],GREEN)
card(s,0.55,5.6,12.23,1.1,"Doble precio + ratios oficiales",
    [[("Cada instrumento guarda su precio en ARS y USD y su ratio de conversión del emisor (Comafi para CEDEAR, SEC F-6/JPMorgan para ADR). El CCL se calcula desde los ADRs y valida el resto.",{"size":10.5})]],
    accent=ORANGE,bg="FDF0E4",hsize=12,head_h=0.32)

# ============================================================ 7 COMO TRABAJAMOS
s=slide(); header(s,"Cómo trabajamos")
cols=[("Código",BLUE,["Rama por historia (feature/S1-07-…)","PR + review de Federico antes de merge","main siempre deployable","CI: lint + tests en cada PR"]),
      ("Seguimiento",GREEN,["Jira: 34 historias · 108 puntos","Mover el estado de cada tarea","Excel de backlog como respaldo","Daily async en el canal"]),
      ("Entregas",ORANGE,["Cada sábado: demo del bloque de la semana","Rama + checklist con referencias","Definición de Hecho por tarea","Sin sorpresas al final"])]
w=3.94
for i,(head_,acc,items) in enumerate(cols):
    x=0.55+i*4.06
    rect(s,x,2.1,w,4.6,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
    rect(s,x,2.1,w,0.55,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08); rect(s,x,2.38,w,0.27,fill=acc)
    text(s,x+0.15,2.1,w-0.3,0.55,head_,size=12.5,color=WHITE,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    text(s,x+0.2,2.8,w-0.36,3.8,bullets(items,acc,size=10),ls=1.16,sa=8)

# ============================================================ 8 ROLES
s=slide(); header(s,"Roles y responsabilidades")
team=[("Federico — Tech Lead",BLUE_D,["Contrato de API + arquitectura","Infra: VPS, Postgres, deploy, backups","Datos → Postgres + scheduler","Code review, QA e integración"]),
      ("Mateo — Backend",BLUE,["FastAPI: endpoints + filtros","Jobs + scheduler","Tests de endpoints + openapi.json","Doble precio / ratios en la API"]),
      ("Joaquín — Frontend Screener",GREEN,["Tabla de 572 empresas + toggle ARS/USD","Filtros + sorting + paginación","Export CSV + panel de jobs","Estados vacío/loading/error + responsive"]),
      ("Aldana — Frontend Detalle",ORANGE,["Ficha de empresa + panel de instrumentos","Dividendos ADR + yield","Vista CEDEAR + comparador","Chips de provenance + disclaimers"])]
w=5.95
for i,(name,acc,items) in enumerate(team):
    x=0.55+(i%2)*6.28; y=2.0+(i//2)*2.5
    rect(s,x,y,w,2.32,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
    rect(s,x,y,0.08,2.32,fill=acc,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.5)
    text(s,x+0.22,y+0.12,w-0.4,0.4,name,size=12.5,color=acc,bold=True)
    text(s,x+0.28,y+0.58,w-0.5,1.65,bullets(items,acc,size=9.4),ls=1.06,sa=3)

# ============================================================ 9-12 SPRINTS (4 columnas por persona)
SP=[("S1","Sprint 1 · Semana 1","Fundamentos navegables — infra + API base + fronts contra los mocks del contrato.",38,GREEN),
    ("S2","Sprint 2 · Semana 2","Datos reales conectados — el front consume la API (572 empresas, doble moneda).",25,BLUE),
    ("S3","Sprint 3 · Semana 3","Funcionalidad completa — filtros, export, jobs en vivo, CEDEAR/ADR, comparador.",26,ORANGE),
    ("S4","Sprint 4 · Semana 4","QA + entrega — tests, hardening, estados/responsive, walkthrough + deploy prod.",19,BLUE_D)]
for sid,title,obj,pts,acc in SP:
    s=slide(); header(s,title,obj+f"   ·   {pts} pts",accent=acc)
    w=2.9
    for i,person in enumerate(PEOPLE):
        x=0.55+i*(w+0.19); pa=ACC[person]
        tasks=[(t,p) for (sp2,t,ep,who,p) in B if sp2==sid and who==person]
        ppts=sum(p for _,p in tasks)
        rect(s,x,2.1,w,4.55,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.04)
        rect(s,x,2.1,w,0.55,fill=pa,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.08); rect(s,x,2.38,w,0.27,fill=pa)
        text(s,x+0.14,2.1,w-0.28,0.55,[[(person+"  ",{"color":WHITE,"bold":True,"size":12}),(f"· {ppts} pts",{"color":WHITE,"size":9.5})]],anchor=MSO_ANCHOR.MIDDLE)
        if tasks:
            body=[[("▸ ",{"color":pa,"bold":True,"size":9.4}),(t+" ",{"color":BODY,"size":9.4}),(f"({p})",{"color":pa,"bold":True,"size":9})] for t,p in tasks]
        else:
            body=[[("—",{"color":GRAYL,"size":10})]]
        text(s,x+0.18,2.8,w-0.34,3.7,body,ls=1.12,sa=7)

# ============================================================ 13 DoD + CONVENCIONES
s=slide(); header(s,"Definición de Hecho + convenciones","Una tarea está 'Hecha' cuando cumple todo esto. Así el sábado no hay sorpresas.")
text(s,0.55,2.05,6,0.35,"Definición de Hecho (DoD)",size=13,color=INK,bold=True)
text(s,0.7,2.5,6,4,bullets([
    "PR aprobado por review + CI en verde",
    "Sin errores en consola",
    "Estados vacío / loading / error cubiertos",
    "Responsive (mobile + desktop)",
    "Datos con provenance / disclaimer donde corresponde",
    "Doc mínima (endpoint o componente)",
],BLUE,size=10.4),ls=1.2,sa=8)
text(s,6.9,2.05,6,0.35,"Convenciones",size=13,color=INK,bold=True)
text(s,7.05,2.5,6,4,bullets([
    "TypeScript estricto en el front",
    "Componentes de shadcn/ui + Tailwind",
    "Nombres de endpoints del contrato, sin desvíos",
    "Formato ARS/USD y flags de calidad consistentes",
    "Commits y ramas con el ID de la historia (S1-07…)",
    "Nada de secretos en el repo (.env)",
],GREEN,size=10.4),ls=1.2,sa=8)

# ============================================================ 14 ESTA SEMANA
s=slide(); header(s,"Esta semana — primeros pasos")
steps=[("01","Accesos","Repo, VPS/staging, Jira y .env para cada uno.",BLUE),
       ("02","Contrato + mocks","Federico congela el contrato de API; mocks disponibles.",GREEN),
       ("03","Scaffolds","FastAPI (Mateo) y Next.js (Joaquín/Aldana) arrancan.",ORANGE),
       ("04","Primera demo","El sábado: infra + API base + fronts navegando contra mocks.",BLUE_D)]
w=2.95
for i,(n,t,d,acc) in enumerate(steps):
    x=0.55+i*3.07
    rect(s,x,2.1,w,2.0,fill=LB3,line=GRID,lw=0.75,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.05)
    text(s,x+0.2,2.25,1.5,0.7,n,size=30,color=acc,bold=True)
    text(s,x+0.2,2.95,w-0.4,0.35,t,size=13,color=INK,bold=True)
    text(s,x+0.2,3.35,w-0.4,0.7,d,size=9.6,color=BODY,ls=1.08)
rect(s,0.55,4.45,12.23,1.55,fill=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,rad=0.03)
items=[("Duración","4 semanas"),("Backlog","34 tareas · 108 pts"),("Cadencia","demo cada sábado"),
       ("Stack","FastAPI + Next.js + Postgres"),("Tracking","Jira + Excel de backlog"),("Arranque","contrato de API día 1")]
for i,(k,v) in enumerate(items):
    x=0.9+(i%3)*4.05; y=4.7+(i//3)*0.62
    text(s,x,y,3.9,0.3,[[(k+"   ",{"color":GRAYL,"size":10}),(v,{"color":WHITE,"size":11,"bold":True})]])
text(s,0.55,6.35,12.2,0.4,"Objetivo de la semana 1: todo navegable contra los mocks del contrato el sábado.",size=10.5,color=GRAY,italic=True,align=PP_ALIGN.CENTER)

import os
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"Kickoff_Equipo_Catalaxia.pptx")
prs.save(out); print("SAVED:",out,"| slides:",len(prs.slides._sldIdLst))
