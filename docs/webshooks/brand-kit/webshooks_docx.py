# -*- coding: utf-8 -*-
"""Sistema de marca webshooks para Word.
Genera un 'reference.docx' con los estilos de webshooks (fuente Calibri, paleta,
encabezados, tablas, footer con webshooks.com + número de página, estilos de
caracter para el wordmark). pandoc lo usa para convertir Markdown -> Word branded.

Uso indirecto: lo llama build_docx.py. Directo:
    python webshooks_docx.py            # crea reference.docx en esta carpeta
"""
import os, sys, glob, subprocess
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---- Paleta de marca webshooks (del PPT v4) ----
NAVY   = RGBColor(0x1A, 0x27, 0x44)
BLUE   = RGBColor(0x3B, 0x82, 0xF6)
ORANGE = RGBColor(0xF5, 0x82, 0x1F)
GREEN  = RGBColor(0x43, 0xA0, 0x47)
BODY   = RGBColor(0x3A, 0x44, 0x58)
GRAY   = RGBColor(0x6B, 0x75, 0x89)
FONT   = "Calibri"

HERE = os.path.dirname(os.path.abspath(__file__))


def find_pandoc():
    from shutil import which
    p = which("pandoc")
    if p:
        return p
    cand = glob.glob(os.path.expandvars(
        r"%LOCALAPPDATA%\Microsoft\WinGet\Packages\JohnMacFarlane.Pandoc_*\pandoc-*\pandoc.exe"))
    return cand[0] if cand else None


def _set_font(style, name=FONT, size=None, color=None, bold=None, italic=None):
    f = style.font
    f.name = name
    # asegurar que aplique a texto latino/complejo
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(a), name)
    if size is not None:  f.size = Pt(size)
    if color is not None: f.color.rgb = color
    if bold is not None:  f.bold = bold
    if italic is not None: f.italic = italic


def _para_style(doc, name, base="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def _char_style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)


def _space(style, before=None, after=None, line=None):
    pf = style.paragraph_format
    if before is not None: pf.space_before = Pt(before)
    if after is not None:  pf.space_after = Pt(after)
    if line is not None:   pf.line_spacing = line


def _footer(doc):
    """Footer: 'webshooks.com' izq + número de página der, con línea superior."""
    sec = doc.sections[0]
    footer = sec.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    # tab derecho al ancho útil (fallback 6.5" si el section no trae medidas)
    try:
        usable = sec.page_width - sec.left_margin - sec.right_margin
    except TypeError:
        usable = None
    if not usable:
        usable = Inches(6.5)
    p.paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    r1 = p.add_run("webshooks.com\t"); r1.font.name = FONT; r1.font.size = Pt(8.5); r1.font.color.rgb = GRAY
    # campo PAGE
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
    rr = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
    sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '17'); rpr.append(sz)   # 8.5pt = 17 half-pts
    col = OxmlElement('w:color'); col.set(qn('w:val'), '6B7589'); rpr.append(col)
    rr.append(rpr)
    t = OxmlElement('w:t'); t.text = '1'; rr.append(t); fld.append(rr)
    p._p.append(fld)
    # línea superior del footer
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    top = OxmlElement('w:top')
    top.set(qn('w:val'), 'single'); top.set(qn('w:sz'), '4')
    top.set(qn('w:space'), '6'); top.set(qn('w:color'), 'D6E2F2')
    pbdr.append(top); pPr.append(pbdr)


def build_reference(out_path=None):
    """Crea el reference.docx branded a partir del default de pandoc."""
    out_path = out_path or os.path.join(HERE, "reference.docx")
    pandoc = find_pandoc()
    if not pandoc:
        sys.exit("No encontré pandoc.")
    # 1) obtener el reference por defecto de pandoc
    r = subprocess.run([pandoc, "--print-default-data-file", "reference.docx"],
                       capture_output=True)
    tmp = os.path.join(HERE, "_default_ref.docx")
    with open(tmp, "wb") as f:
        f.write(r.stdout)

    doc = Document(tmp)

    # 2) estilos base
    _set_font(doc.styles["Normal"], size=10.5, color=BODY)
    for nm, sz, col in [("Heading 1", 17, NAVY), ("Heading 2", 13.5, BLUE),
                        ("Heading 3", 11.5, NAVY), ("Heading 4", 10.5, GRAY)]:
        try:
            st = doc.styles[nm]; _set_font(st, size=sz, color=col, bold=True)
            _space(st, before=(14 if nm == "Heading 1" else 10), after=4)
        except KeyError:
            pass
    for nm, sz, col in [("Title", 30, NAVY), ("Subtitle", 14, BLUE)]:
        try:
            _set_font(doc.styles[nm], size=sz, color=col, bold=(nm == "Title"))
        except KeyError:
            pass
    # tablas: fuente Calibri más chica
    for nm in ("Table", "Table Grid"):
        try:
            _set_font(doc.styles[nm], size=9.5, color=BODY)
        except KeyError:
            pass

    # 3) estilos de párrafo de portada
    def coverpar(name, size, color, before, after, bold=True):
        st = _para_style(doc, name); _set_font(st, size=size, color=color, bold=bold)
        _space(st, before=before, after=after, line=1.0)
    coverpar("WSWordmark", 22, NAVY, 40, 8)
    coverpar("WSCoverTitle", 34, NAVY, 26, 6)
    coverpar("WSCoverSub", 14, BLUE, 0, 26, bold=False)
    coverpar("WSCoverMeta", 11, GRAY, 0, 4, bold=False)
    coverpar("WSLabel", 12, NAVY, 6, 6)

    # 4) estilos de caracter (wordmark + acentos)
    for name, col, bold in [("WSBlue", BLUE, True), ("WSNavy", NAVY, True),
                            ("WSOrange", ORANGE, True), ("WSGreen", GREEN, True),
                            ("WSGray", GRAY, False)]:
        cs = _char_style(doc, name); cs.font.name = FONT
        cs.font.color.rgb = col; cs.font.bold = bold

    # 5) footer
    _footer(doc)

    doc.save(out_path)
    try: os.remove(tmp)
    except OSError: pass
    return out_path


if __name__ == "__main__":
    p = build_reference()
    print("reference.docx ->", p)
