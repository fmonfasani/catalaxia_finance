from __future__ import annotations

import glob
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from docgen.parser import MarkdownParser
from docgen.renderer import MarkdownRenderer
from docgen.theme import ThemeManager


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(str(h).lstrip("#"))


class WordDocumentBuilder:
    """Arma un DOCX branded: portada con logo, índice de Word, footer, y el contenido Markdown."""

    def __init__(self, theme_path: str | Path):
        self.theme_path = Path(theme_path)
        self.theme_manager = ThemeManager(self.theme_path)
        self.parser = MarkdownParser()

    def build(self, pptx_path, markdown_paths, output_path,
              title="Documento", subtitle="", meta="") -> Path:
        theme = self.theme_manager.load_or_extract(pptx_path)
        base = Path(theme.get("_base", self.theme_path.parent))

        doc = Document()
        self._margins(doc)
        self._base_styles(doc, theme)
        self._cover(doc, theme, base, title, subtitle, meta)
        self._index(doc, theme)
        self._footer(doc, theme)

        for i, markdown_path in enumerate(sorted(markdown_paths)):
            if i:
                doc.add_page_break()
            markdown_doc = self.parser.parse_file(markdown_path)
            MarkdownRenderer(doc, theme).render_document(markdown_doc)

        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output)
        return output

    # ---------------- chrome ----------------
    def _margins(self, doc):
        for s in doc.sections:
            s.top_margin = Inches(0.8); s.bottom_margin = Inches(0.8)
            s.left_margin = Inches(0.9); s.right_margin = Inches(0.9)

    def _base_styles(self, doc, t):
        normal = doc.styles["Normal"]
        normal.font.name = t["fonts"]["body"]
        normal.font.size = Pt(t["body"]["size"])
        normal.font.color.rgb = _rgb(t["colors"]["body"])
        for lvl in (1, 2, 3, 4):
            h = t["headings"].get(f"h{lvl}")
            if not h:
                continue
            try:
                st = doc.styles[f"Heading {lvl}"]
                st.font.name = t["fonts"]["title"]
                st.font.size = Pt(h["size"])
                st.font.bold = h.get("bold", True)
                st.font.color.rgb = _rgb(h["color"])
                st.paragraph_format.space_before = Pt(12 if lvl <= 2 else 8)
                st.paragraph_format.space_after = Pt(4)
            except KeyError:
                pass

    def _cover(self, doc, t, base, title, subtitle, meta):
        # logo
        logo = base / t["logo"].get("horizontal", "")
        doc.add_paragraph()
        if logo.exists():
            doc.add_picture(str(logo), width=Inches(2.7))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
        # separación
        for _ in range(4):
            doc.add_paragraph()
        c = t["cover"]
        p = doc.add_paragraph(); r = p.add_run(title)
        r.font.name = t["fonts"]["title"]; r.bold = True
        r.font.size = Pt(c["title"]["size"]); r.font.color.rgb = _rgb(c["title"]["color"])
        # regla naranja
        self._rule(p, t["colors"]["orange"], size=24)
        if subtitle:
            p = doc.add_paragraph(); r = p.add_run(subtitle)
            r.font.name = t["fonts"]["body"]; r.font.size = Pt(c["subtitle"]["size"])
            r.font.color.rgb = _rgb(c["subtitle"]["color"])
        if meta:
            p = doc.add_paragraph(); r = p.add_run(meta)
            r.font.size = Pt(c["meta"]["size"]); r.font.color.rgb = _rgb(c["meta"]["color"])
        doc.add_page_break()

    def _rule(self, paragraph, color, size=18):
        pPr = paragraph._p.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), "6"); bottom.set(qn("w:color"), str(color).lstrip("#"))
        pbdr.append(bottom); pPr.append(pbdr)

    def _index(self, doc, t):
        p = doc.add_paragraph(); r = p.add_run("Índice")
        r.font.name = t["fonts"]["title"]; r.bold = True
        r.font.size = Pt(18); r.font.color.rgb = _rgb(t["colors"]["navy"])
        # campo TOC nativo de Word
        p = doc.add_paragraph()
        run = p.add_run()
        for kind, val in (("begin", None), ("instr", ' TOC \\o "1-2" \\h \\z \\u '),
                          ("separate", None), ("text", "Actualizá el índice: clic derecho → Actualizar campos (F9)."),
                          ("end", None)):
            el = run._r
            if kind == "instr":
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = val
                el.append(it)
            elif kind == "text":
                t2 = OxmlElement("w:t"); t2.text = val; el.append(t2)
            else:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind)
                if kind == "begin":
                    fc.set(qn("w:dirty"), "true")
                el.append(fc)
        doc.add_page_break()

    def _footer(self, doc, t):
        sec = doc.sections[0]
        footer = sec.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0]
        p.text = ""
        usable = sec.page_width - sec.left_margin - sec.right_margin
        p.paragraph_format.tab_stops.add_tab_stop(usable or Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        r = p.add_run(t["footer"].get("text", "webshooks.com") + "\t")
        r.font.name = t["fonts"]["body"]; r.font.size = Pt(8.5); r.font.color.rgb = _rgb(t["colors"]["gray"])
        # campo PAGE
        fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
        rr = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "17"); rpr.append(sz)
        col = OxmlElement("w:color"); col.set(qn("w:val"), t["colors"]["gray"].lstrip("#")); rpr.append(col)
        rr.append(rpr); tt = OxmlElement("w:t"); tt.text = "1"; rr.append(tt); fld.append(rr)
        p._p.append(fld)
        # línea superior
        pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr")
        top = OxmlElement("w:top")
        top.set(qn("w:val"), "single"); top.set(qn("w:sz"), "4")
        top.set(qn("w:space"), "6"); top.set(qn("w:color"), t["footer"].get("rule_color", "#D6E2F2").lstrip("#"))
        pbdr.append(top); pPr.append(pbdr)

    @staticmethod
    def collect_markdown_files(patterns):
        paths = []
        for pattern in patterns:
            for candidate in glob.glob(pattern):
                path = Path(candidate)
                if path.is_file() and path.suffix.lower() == ".md":
                    paths.append(path)
        return sorted(set(paths))
