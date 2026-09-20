# -*- coding: utf-8 -*-
"""Tabla de cobertura de ratios (byma_only + adr) con profundidad realista.
Distingue: valor presente / 'n/d' (no aplica por definicion o economia) / falta (gap).

Reglas de 'no aplica':
  PER, Payout          -> perdida (NI/EPS<=0) o EPS=0 (indefinido)
  P/Book, ROE          -> patrimonio neto < 0
  EV/EBITDA, Deuda/EBITDA -> EBITDA <= 0

Salida: docs/proyecto/Cobertura_Ratios_Catalaxia.xlsx (matriz + resumen) + reporte.
"""
import os, sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY="1A2744"; GREEN_BG="E7F7EC"; GREEN_T="1F6B35"; ND_BG="E9EDF3"; ND_T="6B7589"
MISS_BG="FBE4E4"; MISS_T="C0392B"; WHITE="FFFFFF"; GRID="D6E2F2"; INK="1A2744"; ORANGE="F5821F"
def F(sz=10,b=False,color="3A4458"): return Font(name="Calibri",size=sz,bold=b,color=color)
def fill(c): return PatternFill("solid",fgColor=c)
S=Side(style="thin",color=GRID); BORDER=Border(left=S,right=S,top=S,bottom=S)
CEN=Alignment(horizontal="center",vertical="center"); LEFT=Alignment(horizontal="left",vertical="center")

DB=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","..","data","screener.db")
con=sqlite3.connect(DB); cur=con.cursor()

# ratios a mostrar: (col_screener, nombre, categoria)
RAT=[("ROE","ROE","Rentab."),("ROA","ROA","Rentab."),("MargenNeto","Mg Neto","Rentab."),
     ("margen_operativo","Mg Op.","Rentab."),("DeudaEBITDA","Deuda/EBITDA","Solidez"),
     ("FCF_CE","FCF/CE","Solidez"),("EPS","EPS","Retorno"),("Payout","Payout","Retorno"),
     ("CAGR_EPS_5y","CAGR EPS","Retorno"),("PER","PER","Valuac."),("PriceBook","P/Book","Valuac."),
     ("PriceSales","P/Sales","Valuac."),("ev_ebitda","EV/EBITDA","Valuac.")]
COLS=[r[0] for r in RAT]
sel="cuit,ticker,grupo,EPS,"+",".join(COLS)
rows=cur.execute(f"select {sel} from screener where grupo in ('byma_only','adr') order by grupo desc,ticker").fetchall()

# per_ttm (byma corregido)
pt={r[0]:(r[1],r[2],r[3]) for r in cur.execute("select cuit,estado,per_ttm,ttm_netincome from per_ttm")}
# fundamentales byma: ultimo equity, ebitda
def latest(cuit,concepto):
    r=cur.execute("select valor from cnv_estados_v2 where cuit=? and concepto=? order by period_end desc limit 1",(cuit,concepto)).fetchone()
    return r[0] if r else None
# adr: equity/ebitda desde ratios(edgar) por ticker
adr_f={tk:(eq,eb) for tk,eq,eb in cur.execute("select ticker,_equity,_ebitda_ttm from ratios")}

def num(x):
    return x if isinstance(x,(int,float)) else None

def classify(row):
    cuit,tk,grupo,eps=row[0],row[1],row[2],num(row[3])
    vals={COLS[i]:num(row[4+i]) for i in range(len(COLS))}
    # fundamentales
    if grupo=="byma_only":
        eq=num(latest(cuit,"Equity")); eb=num(latest(cuit,"EBITDA"))
        est,perv,ttm = pt.get(cuit,(None,None,None))
        loss = (num(ttm) is not None and num(ttm)<0) or est=="perdida_real"
        if est is None and eps is not None: loss = eps<0
        per_val = perv if est=="ok" else ("nd" if est=="perdida_real" else None)
    else:
        eq,eb = adr_f.get(tk,(None,None)); eq=num(eq); eb=num(eb)
        loss = (eps is not None and eps<0)
        per_val = vals["PER"]
    neg_eq = eq is not None and eq<0
    neg_eb = eb is not None and eb<=0
    eps_zero = eps==0

    out={}
    def put(key,kind,val=None): out[key]=(kind,val)
    for col,_,_ in RAT:
        v=vals.get(col)
        if col=="PER":
            if per_val=="nd" or loss or eps_zero: put(col,"nd")
            elif isinstance(per_val,(int,float)): put(col,"val",per_val)
            elif isinstance(v,(int,float)) and not loss: put(col,"val",v)
            else: put(col,"miss")
        elif col=="Payout":
            if loss or eps_zero: put(col,"nd")
            elif v is not None: put(col,"val",v)
            else: put(col,"miss")
        elif col in ("PriceBook","ROE"):
            if neg_eq: put(col,"nd")
            elif v is not None: put(col,"val",v)
            else: put(col,"miss")
        elif col in ("ev_ebitda","DeudaEBITDA"):
            if neg_eb: put(col,"nd")
            elif v is not None: put(col,"val",v)
            else: put(col,"miss")
        else:
            if v is not None: put(col,"val",v)
            else: put(col,"miss")
    return out

data=[(r[1],r[2],classify(r)) for r in rows]

# ---------- Excel ----------
wb=Workbook(); ws=wb.active; ws.title="Cobertura"
ws.column_dimensions["A"].width=10; ws.column_dimensions["B"].width=9
for i in range(len(RAT)): ws.column_dimensions[get_column_letter(3+i)].width=11
# header
ws.cell(1,1,"Ticker").fill=fill(NAVY); ws.cell(1,1).font=F(10,True,WHITE); ws.cell(1,1).alignment=CEN; ws.cell(1,1).border=BORDER
ws.cell(1,2,"Grupo").fill=fill(NAVY); ws.cell(1,2).font=F(10,True,WHITE); ws.cell(1,2).alignment=CEN; ws.cell(1,2).border=BORDER
for i,(_,nm,cat) in enumerate(RAT):
    c=ws.cell(1,3+i,nm); c.fill=fill(NAVY); c.font=F(9,True,WHITE); c.alignment=CEN; c.border=BORDER
ws.row_dimensions[1].height=24
r=2
for tk,grupo,cl in data:
    ws.cell(r,1,tk).font=F(9.5,True,INK); ws.cell(r,1).border=BORDER; ws.cell(r,1).alignment=LEFT
    ws.cell(r,2,"BYMA" if grupo=="byma_only" else "ADR").font=F(8.5,False,ND_T); ws.cell(r,2).border=BORDER; ws.cell(r,2).alignment=CEN
    for i,(col,_,_) in enumerate(RAT):
        kind,val=cl[col]; c=ws.cell(r,3+i); c.border=BORDER; c.alignment=CEN; c.font=F(9)
        if kind=="val":
            c.value=round(val,1) if abs(val)>=1 else round(val,2); c.fill=fill(GREEN_BG); c.font=F(9,False,GREEN_T)
        elif kind=="nd":
            c.value="n/d"; c.fill=fill(ND_BG); c.font=F(8.5,False,ND_T)
        else:
            c.value="·"; c.fill=fill(MISS_BG); c.font=F(9,True,MISS_T)
    r+=1
ws.freeze_panes="C2"

# ---------- Resumen ----------
rs=wb.create_sheet("Resumen")
for i,w in enumerate([16,10,8,8,8,14],1): rs.column_dimensions[get_column_letter(i)].width=w
rs.cell(1,1,"COBERTURA EFECTIVA POR RATIO").font=F(13,True,NAVY)
rs.cell(2,1,"efectiva = presentes / (presentes + faltan)  ·  'n/d' se excluye (no aplica)").font=F(9,False,ND_T)
hdr=["Ratio","Categoría","Presente","n/d","Falta","Efectiva"]
for i,h in enumerate(hdr,1):
    c=rs.cell(4,i,h); c.fill=fill(NAVY); c.font=F(9.5,True,WHITE); c.alignment=CEN; c.border=BORDER
r=5
def count(col,grp=None):
    p=nd=m=0
    for tk,grupo,cl in data:
        if grp and grupo!=grp: continue
        k=cl[col][0]
        if k=="val":p+=1
        elif k=="nd":nd+=1
        else:m+=1
    return p,nd,m
for col,nm,cat in RAT:
    p,nd,m=count(col); eff=p/(p+m) if (p+m)>0 else 0
    rs.cell(r,1,nm).font=F(9.5,True,INK); rs.cell(r,2,cat).font=F(9)
    rs.cell(r,3,p).font=F(9.5,False,GREEN_T); rs.cell(r,4,nd).font=F(9.5,False,ND_T); rs.cell(r,5,m).font=F(9.5,False,MISS_T)
    ec=rs.cell(r,6,eff); ec.number_format="0%"; ec.font=F(10,True,ORANGE if eff<0.7 else GREEN_T)
    for i in range(1,7): rs.cell(r,i).border=BORDER; rs.cell(r,i).alignment=CEN if i>1 else LEFT
    r+=1
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),"Cobertura_Ratios_Catalaxia.xlsx")
wb.save(out)

# ---------- reporte a consola ----------
print(f"SAVED: {out}  ({len(data)} activos: 56 BYMA + 17 ADR)")
print("\n== Cobertura efectiva por ratio (n/d excluido) ==")
print(f"{'Ratio':>13} {'pres':>5} {'n/d':>4} {'falta':>6} {'efectiva':>9}   {'bruta(pres/73)':>14}")
for col,nm,cat in RAT:
    p,nd,m=count(col); eff=p/(p+m) if (p+m)>0 else 0; bruta=p/len(data)
    print(f"{nm:>13} {p:>5} {nd:>4} {m:>6} {eff:>8.0%}   {bruta:>13.0%}")
con.close()
