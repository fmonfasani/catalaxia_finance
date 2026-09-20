from pathlib import Path
import hashlib
import json
import csv
import re
import pandas as pd

BASE = Path(__file__).resolve().parents[2]
FIXTURE = BASE / "fixtures" / "tikr_fx_12.csv"
OUTPUT_DIR = BASE / "outputs"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

OUTPUT_JSON = OUTPUT_DIR / "tikr_provenance_audit.json"
OUTPUT_CSV = OUTPUT_DIR / "tikr_provenance_candidates.csv"

MAX_FILE_MB = 50

EXCLUDED_DIRS = {
    ".git",
    "__pycache__",
    "node_modules",
    ".venv",
    "venv",
    "dist",
    "build",
    ".pytest_cache",
}

ALLOWED_TEXT = {
    ".md",
    ".txt",
    ".csv",
    ".json",
    ".yaml",
    ".yml",
}

ALLOWED_EXCEL = {
    ".xlsx",
    ".xlsm",
}


def sha256_file(path):
    h = hashlib.sha256()

    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)

    return h.hexdigest()


def normalize_numeric(value):
    if value is None or isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        try:
            return round(float(value), 2)
        except Exception:
            return None

    text = str(value).strip()

    if not text:
        return None

    text = text.replace(",", "")
    text = re.sub(r"[$€£]", "", text).strip()

    try:
        return round(float(text), 2)
    except Exception:
        return None


def relevant_filename(path):
    name = path.name.lower()

    keywords = (
        "tikr",
        "ypf",
        "quarter",
        "fx",
        "benchmark",
        "financial",
        "ars",
        "usd",
    )

    return any(k in name for k in keywords)


def excluded(path):
    return any(part in EXCLUDED_DIRS for part in path.parts)


if not FIXTURE.exists():
    raise SystemExit(
        f"FAIL: fixture no encontrado: {FIXTURE}"
    )


fixture_sha = sha256_file(FIXTURE)

df = pd.read_csv(FIXTURE)

required = {
    "period",
    "target_date",
    "tikr_fx",
}

missing = required - set(df.columns)

if missing:
    raise SystemExit(
        f"FAIL: fixture missing columns: {sorted(missing)}"
    )


targets = []

for _, row in df.iterrows():

    targets.append(
        {
            "period": str(row["period"]),
            "target_date": str(row["target_date"]),
            "tikr_fx": normalize_numeric(row["tikr_fx"]),
        }
    )


target_values = {
    x["tikr_fx"]
    for x in targets
    if x["tikr_fx"] is not None
}


repo_root = BASE.parents[1]

candidate_files = []

for root in [repo_root, BASE]:

    if not root.exists():
        continue

    for path in root.rglob("*"):

        if not path.is_file():
            continue

        if excluded(path):
            continue

        if path == FIXTURE:
            continue

        suffix = path.suffix.lower()

        if suffix not in ALLOWED_TEXT | ALLOWED_EXCEL:
            continue

        try:
            size_mb = path.stat().st_size / (1024 * 1024)
        except OSError:
            continue

        if size_mb > MAX_FILE_MB:
            continue

        if suffix in ALLOWED_EXCEL:
            if not relevant_filename(path):
                continue

        candidate_files.append(path)


candidate_files = sorted(
    set(candidate_files),
    key=lambda x: str(x).lower()
)


print(f"BASE: {BASE}")
print(f"FIXTURE: {FIXTURE}")
print(f"TARGET VALUES: {sorted(target_values)}")
print(f"CANDIDATE FILES: {len(candidate_files)}")
print()


hits = []


for index, path in enumerate(
    candidate_files,
    start=1,
):

    print(
        f"[{index}/{len(candidate_files)}] "
        f"{path.relative_to(repo_root)}",
        flush=True,
    )

    try:
        file_sha = sha256_file(path)
    except Exception as exc:
        print(
            f"  SHA256 ERROR: {exc}",
            flush=True,
        )
        continue


    suffix = path.suffix.lower()


    if suffix in ALLOWED_TEXT:

        try:
            text = path.read_text(
                encoding="utf-8",
                errors="ignore",
            )
        except Exception:
            continue

        for target in targets:

            value = target["tikr_fx"]

            if value is None:
                continue

            token = f"{value:.2f}"

            if token in text:

                hits.append(
                    {
                        "file": str(path),
                        "file_sha256": file_sha,
                        "format": suffix,
                        "sheet": None,
                        "cell": None,
                        "matched_value": value,
                        "period": target["period"],
                        "target_date": target["target_date"],
                        "match_type": "TEXT_VALUE",
                    }
                )

        continue


    if suffix in ALLOWED_EXCEL:

        try:

            from openpyxl import load_workbook

            wb = load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
            )

        except Exception as exc:

            print(
                f"  XLSX ERROR: {exc}",
                flush=True,
            )

            continue


        try:

            for ws in wb.worksheets:

                for row in ws.iter_rows():

                    for cell in row:

                        value = normalize_numeric(
                            cell.value
                        )

                        if value is None:
                            continue

                        if value not in target_values:
                            continue

                        for target in targets:

                            if target["tikr_fx"] != value:
                                continue

                            hits.append(
                                {
                                    "file": str(path),
                                    "file_sha256": file_sha,
                                    "format": suffix,
                                    "sheet": ws.title,
                                    "cell": cell.coordinate,
                                    "matched_value": value,
                                    "period": target["period"],
                                    "target_date": target["target_date"],
                                    "match_type": "EXCEL_CELL_VALUE",
                                }
                            )

        finally:

            try:
                wb.close()
            except Exception:
                pass


seen = set()
unique_hits = []

for hit in hits:

    key = (
        hit["file"],
        hit["sheet"],
        hit["cell"],
        hit["matched_value"],
        hit["period"],
    )

    if key in seen:
        continue

    seen.add(key)
    unique_hits.append(hit)


hits = unique_hits


matched_values = {
    h["matched_value"]
    for h in hits
}

coverage = (
    len(matched_values) / len(target_values)
    if target_values
    else 0
)


result = {
    "version": "FX-TIKR-PROVENANCE-AUDIT-2.0",
    "fixture": str(FIXTURE),
    "fixture_sha256": fixture_sha,
    "fixture_rows": len(df),
    "target_values": sorted(target_values),
    "candidate_files_scanned": len(candidate_files),
    "candidate_hits": len(hits),
    "matched_distinct_values": len(matched_values),
    "target_distinct_values": len(target_values),
    "coverage": round(coverage, 6),
    "provenance_status": (
        "CANDIDATE_EVIDENCE_FOUND"
        if hits
        else "NO_CANDIDATE_EVIDENCE_FOUND"
    ),
    "decision": "REQUIRES_MANUAL_SOURCE_BINDING",
    "certification_claim": "NONE",
    "hits": hits,
}


OUTPUT_JSON.write_text(
    json.dumps(
        result,
        indent=2,
        ensure_ascii=False,
    ),
    encoding="utf-8",
)


with OUTPUT_CSV.open(
    "w",
    newline="",
    encoding="utf-8",
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "file",
            "file_sha256",
            "format",
            "sheet",
            "cell",
            "matched_value",
            "period",
            "target_date",
            "match_type",
        ],
    )

    writer.writeheader()

    for hit in hits:
        writer.writerow(hit)


print()
print("=" * 70)
print("TIKR PROVENANCE AUDIT — RESULT")
print("=" * 70)
print(f"Fixture rows:            {len(df)}")
print(f"Target FX values:        {len(target_values)}")
print(f"Files scanned:           {len(candidate_files)}")
print(f"Candidate hits:          {len(hits)}")
print(f"Values matched:          {len(matched_values)}")
print(f"Coverage:                {coverage:.2%}")
print(
    "Status:                  "
    f"{result['provenance_status']}"
)
print("Certification claim:     NONE")
print()
print(f"JSON: {OUTPUT_JSON}")
print(f"CSV:  {OUTPUT_CSV}")